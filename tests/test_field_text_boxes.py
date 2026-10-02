"""T-67: надпись из одних полей — не место под содержимое.

Прогон T-66 собрал на корпоративном шаблоне колоду, где тело, таблица и
диаграммы всех контентных слайдов легли в угол слайда. Номер слайда там
нарисован не плейсхолдером номера, а обычной надписью с полем
`<a:fld type="slidenum">`. Правило T-57 «номер слайда и колонтитулы — не
место под содержимое» знало только плейсхолдеры, надпись стала слотом тела,
и место в свободной области макету не строилось.

Шаблон синтетический, из стандартного шаблона python-pptx: тест проверяет
правило, а не подгонку под конкретный файл.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from dpd.models import Bounds, Layout
from dpd.parsing import parse_template

FIELD_BOX = Bounds(x=0.78, y=0.86, w=0.18, h=0.09)
"""Надпись в углу: площадь 0,016 холста — выше порога слота."""

LABEL_BOX = Bounds(x=0.05, y=0.86, w=0.18, h=0.09)
"""Надпись с обычным текстом того же размера: она остаётся слотом."""


def _layout(prs, name: str):
    return next(layout for layout in prs.slide_layouts if layout.name == name)


def _box(prs, target, bounds: Bounds):
    """Надпись на макете: создаётся на черновом слайде и переносится."""
    scratch = prs.slides[0] if len(prs.slides) else prs.slides.add_slide(_layout(prs, "Blank"))
    shape = scratch.shapes.add_textbox(
        int(prs.slide_width * bounds.x),
        int(prs.slide_height * bounds.y),
        int(prs.slide_width * bounds.w),
        int(prs.slide_height * bounds.h),
    )
    target.shapes._spTree.append(shape._element)
    return shape


def _field(shape, kind: str, shown: str) -> None:
    """Абзац надписи — одно поле, как его пишет PowerPoint."""
    paragraph = shape.text_frame.paragraphs[0]._p
    field = etree.SubElement(paragraph, qn("a:fld"), id="{44A0DB5A-4A1A-4420-9295-8BB2A863997E}", type=kind)
    etree.SubElement(field, qn("a:rPr"), lang="ru-RU")
    etree.SubElement(field, qn("a:t")).text = shown


def _build(path: Path) -> Path:
    prs = Presentation()
    _field(_box(prs, _layout(prs, "Title Only"), FIELD_BOX), "slidenum", "‹#›")
    _field(_box(prs, _layout(prs, "Title and Content"), FIELD_BOX), "datetime1", "02.10.2026")
    labelled = _layout(prs, "Section Header")
    _box(prs, labelled, LABEL_BOX).text_frame.text = "Подпись к разделу"
    prs.save(path)
    return path


@pytest.fixture(scope="module")
def layouts(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Layout]:
    path = _build(tmp_path_factory.mktemp("fields") / "synthetic.pptx")
    return {layout.name: layout for layout in parse_template(path, use_cache=False).layouts}


def _overlaps(first: Bounds, second: Bounds) -> bool:
    return (
        first.x < second.x + second.w
        and second.x < first.x + first.w
        and first.y < second.y + second.h
        and second.y < first.y + first.h
    )


def test_slide_number_box_is_not_a_body_slot(layouts: dict[str, Layout]) -> None:
    layout = layouts["Title Only"]
    assert not [slot for slot in layout.slots if slot.origin == "shape"]


def test_layout_with_slide_number_box_gets_room_for_content(layouts: dict[str, Layout]) -> None:
    bodies = [slot for slot in layouts["Title Only"].slots if slot.kind == "body"]
    assert [slot.origin for slot in bodies] == ["derived"]


def test_room_for_content_avoids_slide_number_box(layouts: dict[str, Layout]) -> None:
    derived = next(slot for slot in layouts["Title Only"].slots if slot.origin == "derived")
    assert not _overlaps(derived.bounds, FIELD_BOX)


def test_date_box_is_not_a_body_slot(layouts: dict[str, Layout]) -> None:
    assert not [slot for slot in layouts["Title and Content"].slots if slot.origin == "shape"]


def test_box_with_words_is_still_a_slot(layouts: dict[str, Layout]) -> None:
    shapes = [slot for slot in layouts["Section Header"].slots if slot.origin == "shape"]
    assert len(shapes) == 1
