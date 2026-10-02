"""T-68: разрядка и жирное начертание в оценке вместимости.

Прогон T-66 на корпоративном шаблоне второго трека: заголовок обложки лёг
в полосу макета высотой в одну строку и на рендере ушёл во вторую строку —
она обрезана полосой, — а вёрстка и аудит считали, что он помещается.
Стиль заголовка там жирный и с разрядкой 1,69 pt (`spc="169"` у `defRPr`
списка стилей плейсхолдера). Замер по PDF: строка заголовка в 32 pt
занимает 1044,9 pt при доступных 1039,5; разрядка добавляет к ней около
90 pt, жирные буквы — около 66 pt. Метрика не знала ни того, ни другого и
считала каждую букву обычной шириной 0,55 кегля без разрядки.

Числа ниже — геометрия той обложки (полоса во всю ширину, поля рамки
0,125″ и 0,0625″), сам шаблон в тест не берётся: правило общее, а
материалы второго трека не публикуются.
"""

from __future__ import annotations

import pytest
from lxml import etree
from pptx import Presentation

from dpd.audit.checks.layout import check_text_overflow
from dpd.layout.overflow import plan_compensations, text_fits
from dpd.models import (
    Bounds,
    Canvas,
    Layout,
    RenderedElement,
    RenderedPresentation,
    Slide,
    Slot,
    TextFrame,
    TextRun,
    TextStyle,
)
from dpd.parsing import parse_template

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

CANVAS = Canvas(width_emu=13430250, height_emu=7561263)
STRIP = Bounds(x=0.0, y=3699426 / 7561263, w=1.0, h=692227 / 7561263)
"""Полоса обложки: во всю ширину, высотой в одну строку заголовка."""

PLAIN = TextFrame(
    inset_left=114300 / 13430250,
    inset_right=114300 / 13430250,
    inset_top=57150 / 7561263,
    inset_bottom=57150 / 7561263,
)
SPACED = PLAIN.model_copy(update={"letter_spacing_pt": 1.69})
SPACED_BOLD = SPACED.model_copy(update={"bold": True})

TITLE = ["Масштабирование программы наставничества «Навигатор»"]


def test_letter_spacing_widens_the_line() -> None:
    """В 36 pt заголовок входит в строку только без разрядки."""
    assert text_fits(TITLE, STRIP, CANVAS, 36.0, PLAIN)
    assert not text_fits(TITLE, STRIP, CANVAS, 36.0, SPACED)


def test_bold_letters_are_wider() -> None:
    """В 32 pt разрядки мало — строку переполняют ещё и жирные буквы, как на рендере."""
    assert text_fits(TITLE, STRIP, CANVAS, 32.0, SPACED)
    assert not text_fits(TITLE, STRIP, CANVAS, 32.0, SPACED_BOLD)


def test_negative_spacing_is_honoured_too() -> None:
    """Уплотнённая разрядка сужает строку: метрика честная в обе стороны."""
    narrower = STRIP.model_copy(update={"w": 0.9})
    condensed = PLAIN.model_copy(update={"letter_spacing_pt": -1.5})
    assert not text_fits(TITLE, narrower, CANVAS, 36.0, PLAIN)
    assert text_fits(TITLE, narrower, CANVAS, 36.0, condensed)


def test_layout_picks_a_size_that_fits_with_spacing() -> None:
    """Вёрстка берёт из шкалы кегль, при котором заголовок помещается с разрядкой."""
    slot = Slot(
        id="title-1", kind="title", origin="placeholder", bounds=STRIP, placeholder_idx=0,
        text_style=TextStyle(font="Liberation Sans", size_pt=40.0, resolved_from="layout.lstStyle"),
        frame=SPACED_BOLD,
    )
    compensations, runs = plan_compensations(TITLE, slot, CANVAS, [24.0, 28.0, 32.0, 40.0])
    assert runs[0].size_pt == 28.0
    assert [item.kind for item in compensations] == ["fontScale"]


def test_overflow_check_counts_spacing() -> None:
    """Аудит считает той же метрикой: заголовок в 32 pt за полосой — находка."""
    element = RenderedElement(
        slot_id="title-1", kind="text", bounds=STRIP,
        runs=[TextRun(text=TITLE[0], size_pt=32.0)], frame=SPACED_BOLD,
    )
    deck = RenderedPresentation(
        variant="A", template_hash="sha256:" + "ab" * 32, canvas=CANVAS,
        slides=[Slide(id="s1", layout_id="m/l3", elements=[element])],
    )
    findings = check_text_overflow(deck)
    assert [finding.check_id for finding in findings] == ["layout.text_overflow"]


# --- Разбор кладёт разрядку и начертание в рамку слота ----------------------


def _layout(prs, name: str):
    return next(layout for layout in prs.slide_layouts if layout.name == name)


def _title(shapes):
    return next(ph for ph in shapes.placeholders if ph.placeholder_format.idx == 0)


def _first_level_run(placeholder):
    """`defRPr` первого уровня списка стилей плейсхолдера, созданный при нужде."""
    body = placeholder._element.find(P + "txBody")
    list_style = body.find(A + "lstStyle")
    if list_style is None:
        list_style = etree.Element(A + "lstStyle")
        body.find(A + "bodyPr").addnext(list_style)
    level = list_style.find(A + "lvl1pPr")
    if level is None:
        level = etree.SubElement(list_style, A + "lvl1pPr")
    run = level.find(A + "defRPr")
    if run is None:
        run = etree.SubElement(level, A + "defRPr")
    return run


@pytest.fixture(scope="module")
def spaced(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Layout]:
    prs = Presentation()

    # Разрядка и жирный — у плейсхолдера заголовка макета.
    own = _first_level_run(_title(_layout(prs, "Title Slide")))
    own.set("spc", "169")
    own.set("b", "1")

    # Жирный — в стиле заголовков мастера: макеты без своего наследуют его.
    styles = prs.slide_master.element.find(P + "txStyles").find(P + "titleStyle")
    styles.find(A + "lvl1pPr").find(A + "defRPr").set("b", "1")

    # Макет отменяет жирный мастера своим `b="0"`.
    _first_level_run(_title(_layout(prs, "Section Header"))).set("b", "0")

    path = tmp_path_factory.mktemp("spacing") / "spaced.pptx"
    prs.save(path)
    return {layout.name: layout for layout in parse_template(path, use_cache=False).layouts}


def _title_frame(layout: Layout) -> TextFrame:
    frame = next(slot for slot in layout.slots if slot.kind == "title").frame
    assert frame is not None
    return frame


def test_parser_reads_spacing_and_weight_of_the_placeholder(spaced: dict[str, Layout]) -> None:
    frame = _title_frame(spaced["Title Slide"])
    assert frame.letter_spacing_pt == pytest.approx(1.69)
    assert frame.bold is True


def test_parser_inherits_the_weight_from_the_master(spaced: dict[str, Layout]) -> None:
    frame = _title_frame(spaced["Title Only"])
    assert frame.letter_spacing_pt == 0.0
    assert frame.bold is True


def test_layout_overrides_the_weight_of_the_master(spaced: dict[str, Layout]) -> None:
    assert _title_frame(spaced["Section Header"]).bold is False


def test_body_without_weight_stays_regular(spaced: dict[str, Layout]) -> None:
    body = next(slot for slot in spaced["Title and Content"].slots if slot.kind == "body")
    assert body.frame is not None
    assert body.frame.bold is False
    assert body.frame.letter_spacing_pt == 0.0
