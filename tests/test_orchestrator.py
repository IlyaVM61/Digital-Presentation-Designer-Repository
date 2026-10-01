"""T-38: последовательность этапов прогона.

Интерфейсу нужно что вызывать: до сих пор слои складывались вручную в тестах
и демонстрациях. Оркестратор собирает их в один путь — от файла шаблона до
выгруженной колоды, — и отдаёт наружу две вещи, без которых интерфейс
бесполезен: события прогресса и время каждого этапа.

**Время замеряется здесь, а не в интерфейсе.** Бюджет ТЗ — пять минут на
полный цикл (решение D4), и разбивка по этапам нужна, чтобы понимать, что
именно не уложилось. Интерфейс её только показывает.

**LibreOffice запускается один раз.** Он же самая дорогая операция пайплайна:
13–28 с на колоду. Превью и PDF нуждаются в одной и той же конвертации, и
второй запуск стоил бы столько же — это проверяется отдельно.
"""

from __future__ import annotations

import os
from collections.abc import Iterator
from pathlib import Path

import pytest
import yaml

from dpd.generation import structure_from_outline
from dpd.orchestrator import STAGES, run_pipeline
from dpd.render import soffice_path

CALIBRATION = Path(__file__).resolve().parents[1] / "assets" / "templates" / "calibration"
TEMPLATE = "VK Tech шаблон.pptx"

OUTLINE = """Итоги пилота
Что изменилось за квартал
- Срок сборки колоды сократился с девяти дней до трёх
- Правки текста переживают сохранение файла
Как шли
- Разбор шаблона
- Вёрстка
- Аудит
"""


@pytest.fixture(scope="module")
def out_dir() -> Iterator[Path]:
    base = Path(os.environ.get("DPD_OUTPUT_DIR", "D:/dpd-out"))
    path = base / "tests" / "orchestrator"
    path.mkdir(parents=True, exist_ok=True)
    yield path


@pytest.fixture(scope="module")
def template() -> Path:
    path = CALIBRATION / TEMPLATE
    if not path.exists():
        pytest.skip(f"нет калибровочного шаблона {TEMPLATE}")
    return path


# --- Сквозной путь ---------------------------------------------------------


def test_run_goes_from_template_file_to_exported_deck(template: Path, out_dir: Path) -> None:
    """Критерий приёмки задачи: путь проходится целиком.

    Формат `.pptx` берётся один: он обязателен, быстр и именно его ТЗ
    требует собирать нативными объектами.
    """
    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "run", formats=("pptx",))

    assert len(result.variants) == 3, "три варианта — требование ТЗ"
    assert result.exports["pptx"].is_file()
    assert result.chosen.variant == result.variants[0].variant
    assert result.template.layouts, "схема шаблона не вернулась"


def test_every_stage_is_timed(template: Path, out_dir: Path) -> None:
    """Разбивка по этапам — то, ради чего считается бюджет в пять минут."""
    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "timings", formats=("pptx",))

    assert set(result.timings) == {"parse", "layout", "audit", "export"}
    assert all(value >= 0 for value in result.timings.values())
    assert result.seconds >= sum(result.timings.values()) * 0.9


def test_progress_is_reported_stage_by_stage(template: Path, out_dir: Path) -> None:
    """Интерфейсу нужно показывать, что происходит: прогон идёт десятки секунд."""
    events: list[tuple[str, int, int]] = []

    run_pipeline(
        template,
        structure_from_outline(OUTLINE),
        out_dir / "progress",
        formats=("pptx",),
        progress=lambda stage, done, total: events.append((stage, done, total)),
    )

    # Этап шлёт не одно событие, а несколько: начало, шаги, конец. Важно,
    # что этапы идут в каноническом порядке и не перемежаются друг другом.
    порядок = [
        name
        for index, (name, _, _) in enumerate(events)
        if index == 0 or events[index - 1][0] != name
    ]
    случившиеся = {name for name, _, _ in events}
    assert порядок == [stage for stage in STAGES if stage in случившиеся]
    assert events[-1][1] == events[-1][2], "последнее событие должно сообщать о завершении"


def test_audit_runs_for_every_variant(template: Path, out_dir: Path) -> None:
    """Аудит формы выполняется для всех трёх вариантов: вёрстка у них разная."""
    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "audit", formats=("pptx",))

    assert set(result.reports) == {variant.variant for variant in result.variants}
    for report in result.reports.values():
        assert not [item for item in report.findings if item.fixability == "mechanical"], (
            "механические находки должны быть исправлены починкой"
        )


# --- Форматы ---------------------------------------------------------------


def test_html_export_needs_no_libreoffice(template: Path, out_dir: Path) -> None:
    """HTML собирается из контракта: тяжёлой конвертации для него не нужно."""
    result = run_pipeline(
        template, structure_from_outline(OUTLINE), out_dir / "html", formats=("pptx", "html")
    )

    assert result.exports["html"].is_file()
    assert "Итоги пилота" in result.exports["html"].read_text(encoding="utf-8")


def test_only_requested_formats_are_offered(template: Path, out_dir: Path) -> None:
    """Выгружается ровно то, что попросили.

    `.pptx` собирается всегда — он основа колоды и источник для конвертации
    в PDF, — но попадать в выдачу он не должен, если его не просили.
    Пользователь снял галочку PowerPoint и всё равно получал файл: интерфейс
    обещал выбор, которого не было.
    """
    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "formats", formats=("html",))

    assert set(result.exports) == {"html"}


def test_libreoffice_is_started_once_for_previews_and_pdf(
    template: Path, out_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Превью и PDF всех трёх вариантов довольствуются одной конвертацией.

    Это самая дорогая операция пайплайна — 13–28 с, — и второй запуск стоил
    бы столько же на тот же результат. При бюджете в пять минут на полный
    цикл такая трата недопустима. С T-39 в конвертацию идут все три
    варианта: их показывают рядом, и колоды передаются одной пачкой.
    """
    if soffice_path() is None:
        pytest.skip("LibreOffice не найден")

    from dpd import orchestrator

    calls: list[list[Path]] = []
    original = orchestrator.convert_to_pdfs

    def counted(pptx_paths, out, **kwargs):
        calls.append([Path(path) for path in pptx_paths])
        return original(pptx_paths, out, **kwargs)

    monkeypatch.setattr(orchestrator, "convert_to_pdfs", counted)

    result = run_pipeline(
        template,
        structure_from_outline(OUTLINE),
        out_dir / "once",
        formats=("pptx", "pdf"),
        previews=True,
    )

    assert len(calls) == 1, f"LibreOffice запущен {len(calls)} раз"
    assert len(calls[0]) == 3, "в конвертацию попали не все варианты"
    assert result.exports["pdf"].is_file()
    assert len(result.previews) == len(result.chosen.slides)
    assert result.timings["render"] > 0

    # T-39: превью каждого варианта свои, и в одну папку они не сваливаются.
    assert set(result.variant_previews) == {deck.variant for deck in result.variants}
    for deck in result.variants:
        images = result.variant_previews[deck.variant]
        assert len(images) == len(deck.slides), f"вариант {deck.variant}: превью не на каждый слайд"
    assert len({image for images in result.variant_previews.values() for image in images}) == 3 * len(
        result.chosen.slides
    ), "превью одного варианта перезаписали превью другого"


# --- Выбор варианта --------------------------------------------------------


def test_chosen_variant_can_be_named(template: Path, out_dir: Path) -> None:
    """Пользователь выбирает вариант — экспорт идёт для выбранного."""
    result = run_pipeline(
        template, structure_from_outline(OUTLINE), out_dir / "choice", formats=("pptx",), variant="C"
    )

    assert result.chosen.variant == "C"


def test_unknown_variant_is_a_loud_error(template: Path, out_dir: Path) -> None:
    with pytest.raises(ValueError, match="вариант"):
        run_pipeline(
            template, structure_from_outline(OUTLINE), out_dir / "bad", formats=("pptx",), variant="Я"
        )


# --- T-39: три варианта рядом и выбор после прогона -------------------------
#
# Пользователь выбирает вариант, посмотрев на все три. Значит, к концу прогона
# готовы файлы каждого: выбор не должен стоить второго прогона и второго
# запуска LibreOffice. Цена измерена: +1–2 с на `.pptx` двух лишних
# вариантов, +8 с на их конвертацию, HTML почти бесплатен.


def test_every_variant_is_exported_for_choice_after_the_run(template: Path, out_dir: Path) -> None:
    result = run_pipeline(
        template, structure_from_outline(OUTLINE), out_dir / "all", formats=("pptx", "html")
    )

    assert set(result.variant_exports) == {"A", "B", "C"}
    for variant, files in result.variant_exports.items():
        assert set(files) == {"pptx", "html"}, f"вариант {variant}: не все форматы"
        assert all(path.is_file() for path in files.values())
        assert all(f"-{variant}." in path.name for path in files.values()), "имя файла не называет вариант"
    assert result.exports == result.variant_exports[result.chosen.variant]


def test_variant_names_come_from_the_config(template: Path, out_dir: Path) -> None:
    """Интерфейс называет варианты по-человечески, а имена живут в конфиге."""
    from dpd.layout.variants import load_profiles

    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "names", formats=("html",))

    assert result.variant_names == {profile.id: profile.name for profile in load_profiles()}


def test_identical_variants_are_reported(
    template: Path, out_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Три одинаковые колоды нельзя молча выдать за три варианта.

    Стратегия вариантов прямо требует предъявить этот факт пользователю: на
    бедном шаблоне оси регистра могут не сработать. Одинаковые профили
    воспроизводят такой шаблон без подбора файла.
    """
    same = {"prefer_scheme": "any", "prefer_families": ["content"], "size_shift": 0, "layout_offset": 0}
    config = tmp_path / "variants.yaml"
    config.write_text(
        yaml.safe_dump(
            {"version": 1, "profiles": [{"id": name, "name": f"Вариант {name}", **same} for name in "ABC"]},
            allow_unicode=True,
        ),
        encoding="utf-8",
    )
    monkeypatch.setenv("DPD_VARIANTS_CONFIG", str(config))

    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "same", formats=("html",))

    assert [finding.check_id for finding in result.distinction] == ["variants.low_distinction"]


def test_distinct_variants_raise_no_alarm(template: Path, out_dir: Path) -> None:
    """На нормальном шаблоне проверка молчит — иначе предупреждению не поверят."""
    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "distinct", formats=("html",))

    assert result.distinction == []


# --- T-40: работа с находками ------------------------------------------------
#
# Выбор исправления стоит одного шага, а не второго прогона: пользователь уже
# смотрит на варианты, и пересобирать все три ради одного слайда незачем.
# Но файлы выбранного варианта обязаны измениться — исправление, которого нет
# в скачанном файле, пользователю ничего не дало.

CROWDED = """Итоги пилота
Что изменилось за квартал
- Срок сборки колоды сократился с девяти дней до трёх
- Правки текста переживают сохранение файла
Что умеет система
- Разбор шаблона занимает меньше секунды
- Три варианта собираются одним нажатием
- Правки текста переживают сохранение
- Таблицы остаются таблицами
- Диаграммы остаются диаграммами
- Проверка оформления идёт до выгрузки
- Механические дефекты исправляются сами
- Файлы готовы в трёх форматах
"""


def test_finding_fix_reaudit_cycle_goes_through(template: Path, out_dir: Path) -> None:
    """Критерий приёмки T-40: находка → исправление → повторный аудит → файлы."""
    from pptx import Presentation

    from dpd.orchestrator import revise

    result = run_pipeline(template, structure_from_outline(CROWDED), out_dir / "revise", formats=("pptx", "html"))
    variant = result.variants[0].variant
    report = result.reports[variant]
    index = next(i for i, f in enumerate(report.findings) if f.check_id == "density.too_many_bullets")
    assert [item.id for item in result.remedies[variant][index]] == ["split"]
    slides = len(Presentation(result.variant_exports[variant]["pptx"]).slides)
    autofixed = [f for f in report.findings if f.status == "autofixed"]
    untouched = {key: value for key, value in result.reports.items() if key != variant}

    revised = revise(result, variant, {index: "split"})

    after = revised.reports[variant]
    assert "density.too_many_bullets" not in {f.check_id for f in after.findings}
    assert len(Presentation(revised.variant_exports[variant]["pptx"]).slides) == slides + 1
    html = revised.variant_exports[variant]["html"].read_text(encoding="utf-8")
    assert html.count('<section class="slide"') == slides + 1
    assert "Файлы готовы в трёх форматах" in html
    assert [item.remedy.id for item in revised.revision.applied] == ["split"]
    assert revised.revision.before.get("warning", 0) > revised.revision.after.get("warning", 0)
    assert [f for f in after.findings if f.status == "autofixed"][: len(autofixed)] == autofixed, (
        "автоисправления первого прогона пропали из отчёта"
    )
    assert {key: revised.reports[key] for key in untouched} == untouched, "задеты другие варианты"
    assert revised.chosen.variant == result.chosen.variant
    assert len(revised.chosen.slides) == slides + 1


def test_revision_without_choices_changes_nothing(template: Path, out_dir: Path) -> None:
    from dpd.orchestrator import revise

    result = run_pipeline(template, structure_from_outline(CROWDED), out_dir / "keep", formats=("html",))
    variant = result.variants[1].variant

    revised = revise(result, variant, {})

    assert revised.reports[variant] == result.reports[variant]
    assert not revised.revision.applied


# --- T-41: разбивка времени в отчёте ----------------------------------------
#
# Бюджет ТЗ — пять минут на полный цикл, считая разбор шаблона (решение D4).
# Разбивка по этапам закрывает вопрос о трактовке при любом прочтении, но
# только если она лежит в отчёте, а не живёт в памяти до конца прогона
# (NFR-1а).


def test_every_report_carries_the_run_timings(template: Path, out_dir: Path) -> None:
    """Критерий приёмки T-41: `timings` отчёта заполнен.

    Этапы общие для трёх вариантов — разбор один, конвертация одна, — поэтому
    каждый отчёт несёт разбивку прогона целиком. Делить время по вариантам
    значило бы его выдумывать.
    """
    result = run_pipeline(
        template, structure_from_outline(OUTLINE), out_dir / "report-timings", formats=("pptx", "html")
    )

    assert set(result.timings) == {"parse", "layout", "audit", "export"}
    for variant, report in result.reports.items():
        assert report.timings == result.timings, f"вариант {variant}: в отчёте нет разбивки времени"


def test_timings_follow_the_stage_order(template: Path, out_dir: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Разбивка читается сверху вниз как ход прогона.

    Выгрузка `.pptx` начинается раньше рендера, и словарь, заполняемый по
    ходу, ставил бы «Сохраняем файлы» перед «Готовим превью». LibreOffice
    подменён: проверяется порядок, а не конвертация.
    """
    from dpd import orchestrator

    monkeypatch.setattr(
        orchestrator, "convert_to_pdfs", lambda paths, out, **_: [Path(path).with_suffix(".pdf") for path in paths]
    )
    monkeypatch.setattr(orchestrator, "render_pdf_pages", lambda pdf, out, **_: [])

    result = run_pipeline(
        template, structure_from_outline(OUTLINE), out_dir / "order", formats=("pptx",), previews=True
    )

    assert list(result.timings) == list(STAGES)
    assert list(result.reports[result.chosen.variant].timings) == list(STAGES)


def test_revised_report_keeps_the_run_timings(template: Path, out_dir: Path) -> None:
    """Исправление по выбору пересобирает отчёт, но разбивку не стирает.

    Время самого исправления в бюджет цикла не входит: между прогоном и
    исправлением лежит решение пользователя, и оно занимает сколько угодно.
    """
    from dpd.orchestrator import revise

    result = run_pipeline(template, structure_from_outline(CROWDED), out_dir / "revise-timings", formats=("html",))
    variant = result.variants[0].variant
    report = result.reports[variant]
    index = next(i for i, f in enumerate(report.findings) if f.check_id == "density.too_many_bullets")

    revised = revise(result, variant, {index: "split"})

    assert revised.reports[variant].findings != report.findings, "исправление не состоялось"
    assert revised.reports[variant].timings == result.timings


# --- T-48: версии промптов в отчёте -----------------------------------------
#
# Какими промптами получен результат, должно быть видно из отчёта прогона:
# без этого воспроизводимость недетерминированной части непроверяема (ТЗ,
# п. 2.4). Промпты загружаются на старте прогона, и сломанный файл
# останавливает его до разбора шаблона, а не посреди генерации.


def test_every_report_names_the_prompt_set_and_revision_keeps_it(template: Path, out_dir: Path) -> None:
    from dpd.llm import load_prompts
    from dpd.orchestrator import revise

    expected = load_prompts().versions()
    result = run_pipeline(template, structure_from_outline(CROWDED), out_dir / "prompts", formats=("html",))

    assert {key: report.prompt_versions for key, report in result.reports.items()} == {
        key: expected for key in result.reports
    }

    variant = result.variants[0].variant
    report = result.reports[variant]
    index = next(i for i, f in enumerate(report.findings) if f.check_id == "density.too_many_bullets")
    revised = revise(result, variant, {index: "split"})

    assert revised.reports[variant].prompt_versions == expected


# --- Аудит текста (T-51) -----------------------------------------------------


def test_text_audit_reaches_every_variant_and_survives_revision(template: Path, out_dir: Path) -> None:
    """Аудит текста выполняется один раз на колоду, до вёрстки, а его находки
    видны в отчёте каждого варианта — и после исправления по выбору
    пользователя тоже, на своём слайде, хотя номера сдвинулись."""
    from dpd.audit.textual import audit_text
    from dpd.models import Finding
    from dpd.orchestrator import revise

    structure = structure_from_outline(CROWDED + "Что дальше\n- Сборка по контент-пакету\n")
    last = structure.slides[-1]
    typo = Finding.model_validate(
        {
            "checkId": "content.typos",
            "category": "content",
            "class": "model",
            "severity": "warning",
            "fixability": "semantic",
            "message": "Опечатка",
            "slideNumber": len(structure.slides),
            "evidence": {"slideId": last.id},
        }
    )
    text = audit_text(structure)
    text = text.model_copy(update={"findings": [typo], "checks_run": [*text.checks_run, "content.typos"]})

    result = run_pipeline(template, structure, out_dir / "text-audit", formats=("html",), text_audit=text)

    for report in result.reports.values():
        assert "content.typos" in [f.check_id for f in report.findings]
        assert "content.headline_no_conclusion" in report.checks_skipped
    variant = result.variants[0].variant
    report = result.reports[variant]
    index = next(i for i, f in enumerate(report.findings) if f.check_id == "density.too_many_bullets")

    revised = revise(result, variant, {index: "split"})

    [moved] = [f for f in revised.reports[variant].findings if f.check_id == "content.typos"]
    assert moved.slide_number == len(structure.slides) + 1


# --- Аудит визуала выбранного варианта (T-52) --------------------------------

VISUAL = ("content.irrelevant_imagery", "content.visual_readability")


def fake_render(monkeypatch: pytest.MonkeyPatch) -> list[Path]:
    """LibreOffice подменён: проверяется, кто и что смотрит, а не конвертация.

    Превью по одному на слайд — их число берётся из колоды, лежащей рядом с
    «PDF», как у настоящего рендера.
    """
    from pptx import Presentation

    from dpd import orchestrator

    converted: list[Path] = []

    def convert(paths, out, **_):
        converted.extend(Path(path) for path in paths)
        return [Path(path).with_suffix(".pdf") for path in paths]

    def render(pdf, out, **_):
        out = Path(out)
        out.mkdir(parents=True, exist_ok=True)
        count = len(Presentation(str(Path(pdf).with_suffix(".pptx"))).slides)
        images = [out / f"slide-{number:03d}.png" for number in range(1, count + 1)]
        for image in images:
            image.write_bytes(b"png")
        return images

    monkeypatch.setattr(orchestrator, "convert_to_pdfs", convert)
    monkeypatch.setattr(orchestrator, "render_pdf_pages", render)
    return converted


def vision(requests: list[dict]):
    """VLM-подделка: на слайдах не видно ни строки."""
    import json

    import httpx

    from dpd.llm import ModelClient, ModelSettings, load_prompts

    def reply(request: httpx.Request) -> httpx.Response:
        requests.append(json.loads(request.content))
        content = {"lines": [], "offTopicPictures": []}
        message = {"role": "assistant", "content": json.dumps(content, ensure_ascii=False)}
        return httpx.Response(200, json={"choices": [{"message": message, "finish_reason": "stop"}]})

    settings = ModelSettings(
        base_url="https://provider.test/v1",
        api_key="test",
        model="test-vlm",
        temperature=0,
        seed=1,
        max_tokens=100,
        timeout_sec=5,
        max_attempts=1,
        retry_delay_sec=0,
    )
    prompts = load_prompts()
    return ModelClient(settings, repair_prompt=prompts.get("skills/response-repair").text, transport=httpx.MockTransport(reply)), prompts


def test_visual_checks_are_named_skipped_until_a_variant_is_looked_at(template: Path, out_dir: Path) -> None:
    """Прогон модели не зовёт: смотрится выбранный вариант, а выбирают после
    прогона. До этого отчёт называет проверки визуала пропущенными, а не молчит."""
    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "visual-skipped", formats=("html",))

    for report in result.reports.values():
        assert set(VISUAL) <= set(report.checks_skipped)
        assert not set(VISUAL) & set(report.checks_run)


def test_look_audits_the_chosen_variant_only(
    template: Path, out_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Критерий T-52: один запрос на слайд, все вопросы разом — и только для
    выбранного варианта (ADR-0002, п. 4). Превью прогона переиспользуются:
    второй запуск LibreOffice ради того же изображения не нужен."""
    from dpd.orchestrator import look

    converted = fake_render(monkeypatch)
    result = run_pipeline(
        template, structure_from_outline(OUTLINE), out_dir / "visual-look", formats=("pptx",), previews=True
    )
    launches = len(converted)
    requests: list[dict] = []
    variant = result.variants[1].variant
    deck = next(item for item in result.variants if item.variant == variant)

    looked = look(result, variant, *vision(requests))

    assert len(requests) == len(deck.slides)
    assert len(converted) == launches
    report = looked.reports[variant]
    assert set(VISUAL) <= set(report.checks_run)
    assert sorted({f.slide_number for f in report.findings if f.check_id == "content.visual_readability"}) == list(
        range(1, len(deck.slides) + 1)
    )
    assert "visual" in report.timings
    for other, unseen in looked.reports.items():
        if other != variant:
            assert set(VISUAL) <= set(unseen.checks_skipped)


def test_look_renders_the_variant_when_the_run_had_no_previews(
    template: Path, out_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from dpd.orchestrator import look

    converted = fake_render(monkeypatch)
    result = run_pipeline(template, structure_from_outline(OUTLINE), out_dir / "visual-render", formats=("html",))
    assert converted == []
    requests: list[dict] = []
    variant = result.chosen.variant

    looked = look(result, variant, *vision(requests))

    assert len(converted) == 1
    assert len(looked.variant_previews[variant]) == len(looked.chosen.slides) == len(requests)


def test_revision_drops_what_the_model_saw_before_it(
    template: Path, out_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """После исправления слайды другие, а модель смотрела на прежние: её
    замечания не переносятся, проверки снова названы пропущенными."""
    from dpd.orchestrator import look, revise

    fake_render(monkeypatch)
    result = run_pipeline(template, structure_from_outline(CROWDED), out_dir / "visual-revise", formats=("html",))
    variant = result.variants[0].variant
    looked = look(result, variant, *vision([]))
    report = looked.reports[variant]
    index = next(i for i, f in enumerate(report.findings) if f.check_id == "density.too_many_bullets")

    revised = revise(looked, variant, {index: "split"})

    fresh = revised.reports[variant]
    assert not [f for f in fresh.findings if f.check_id in VISUAL]
    assert set(VISUAL) <= set(fresh.checks_skipped)


# --- Черновик колоды по контент-пакету (T-58) ---------------------------------
#
# Интерфейс и батч пишут колоду одинаково: структура и содержание от модели,
# затем аудит текста — до вёрстки. Страница ничего не считает сама, поэтому
# этот шаг живёт здесь, рядом с прогоном, с событиями прогресса и замером.

PACK = Path(__file__).resolve().parents[1] / "assets" / "content-pack"


def writer(fail_on: str | None = None):
    """LLM-подделка батча T-50; `fail_on` — схема ответа, на которой провайдер отказывает."""
    import httpx
    from test_batch import SETTINGS, fake_model

    from dpd.llm import build_client, load_prompts

    def reply(request: httpx.Request) -> httpx.Response:
        import json

        if json.loads(request.content)["response_format"]["json_schema"]["name"] == fail_on:
            return httpx.Response(400, json={"error": "отказ провайдера"})
        return fake_model(request)

    prompts = load_prompts()
    return build_client(SETTINGS, prompts, transport=httpx.MockTransport(reply)), prompts


def test_draft_writes_the_deck_and_checks_its_text() -> None:
    """Критерий T-58 со стороны оркестратора: колода от модели и аудит её
    текста одним вызовом, со временем каждого шага — оно входит в бюджет
    полного цикла (D4)."""
    from test_batch import GENERATED_BODY

    from dpd.generation import load_content_pack
    from dpd.orchestrator import DRAFT_STAGES, draft

    events: list[str] = []
    drafted = draft(load_content_pack(PACK), *writer(), progress=lambda stage, done, total: events.append(stage))

    assert any(slide.body and GENERATED_BODY in slide.body.items for slide in drafted.structure.slides)
    assert "content.headline_no_conclusion" in drafted.text_audit.checks_run
    assert not [item for item in drafted.text_audit.checks_skipped if item.endswith(":model")]
    assert drafted.review_error is None
    assert list(drafted.text_audit.timings) == list(DRAFT_STAGES)
    assert list(dict.fromkeys(events)) == list(DRAFT_STAGES)


def test_failed_text_check_does_not_stop_the_deck() -> None:
    """Сбой модели на проверке текста колоды не останавливает — как в батче:
    смысловые проверки названы пропущенными, причина возвращается, чтобы
    интерфейс её показал, а не проглотил."""
    from dpd.generation import load_content_pack
    from dpd.orchestrator import draft

    drafted = draft(load_content_pack(PACK), *writer(fail_on="TextReview"))

    assert drafted.structure.slides
    assert drafted.review_error and "HTTP 400" in drafted.review_error
    assert [item for item in drafted.text_audit.checks_skipped if item.endswith(":model")]


def test_failed_generation_stops_before_the_layout() -> None:
    """Частичной колоды дальше не уходит (EF-8): сбой генерации — исключение."""
    from dpd.generation import load_content_pack
    from dpd.llm import ModelError
    from dpd.orchestrator import draft

    with pytest.raises(ModelError):
        draft(load_content_pack(PACK), *writer(fail_on="DeckOutline"))


def test_draft_time_leads_the_run_timings(template: Path, out_dir: Path) -> None:
    """Разбивка времени принадлежит отчёту (NFR-1а): время написания и
    проверки текста стоит в нём первым, по ходу цикла, а не теряется в
    интерфейсе."""
    from dpd.generation import load_content_pack
    from dpd.orchestrator import DRAFT_STAGES, draft

    drafted = draft(load_content_pack(PACK), *writer())
    result = run_pipeline(template, drafted.structure, out_dir / "draft-timings", formats=("html",), text_audit=drafted.text_audit)

    stages = [stage for stage in STAGES if stage in result.timings]
    for report in result.reports.values():
        assert list(report.timings) == [*DRAFT_STAGES, *stages]
