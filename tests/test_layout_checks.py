"""T-28: семь проверок вёрстки класса `file`.

Критерий приёмки задачи: **на слайде с заведомым дефектом каждая проверка
срабатывает, на чистом — нет**. Обе половины обязательны. Проверка, которая
молчит всегда, не отличается от отсутствующей; проверка, которая срабатывает
всегда, обесценивается за один прогон — и это опаснее, потому что выглядит
работой.

Дефекты собираются синтетически: вёрстка кладёт содержимое ровно в слоты
шаблона и сама таких слайдов не порождает. Поэтому в конце файла стоит
прогон на калибровочных шаблонах — он проверяет вторую половину критерия на
живых данных, а не на выдуманных.

Разделение трёх похожих проверок задано Приложением 1 и здесь закреплено
тестами: `out_of_bounds` — про рамку элемента и край холста, `text_overflow` —
про текст и его собственную рамку, `text_clipped` — про текст, который
край холста действительно срезал.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from dpd.audit import REGISTRY, AuditContext, discover, run_checks
from dpd.layout import compose_variants
from dpd.models import (
    Bounds,
    Canvas,
    Layout,
    PresentationStructure,
    RenderedElement,
    RenderedPresentation,
    Slide,
    SlideBody,
    Slot,
    StructureSlide,
    TemplateSchema,
    TemplateSource,
    TextRun,
)
from dpd.parsing import parse_template

CALIBRATION = Path(__file__).resolve().parents[1] / "assets" / "templates" / "calibration"
CANVAS = Canvas(width_emu=12192000, height_emu=6858000)

# Классификация из `docs/07-testing/audit-checklist.md`. Реестр и чек-лист
# обязаны совпадать: чек-лист — нормативный документ сдачи, и расхождение с
# ним означает, что один из двух врёт.
CHECKLIST = {
    "layout.out_of_bounds": ("file", "critical", "mechanical"),
    "layout.overlap": ("file", "critical", "lossy"),
    "layout.text_overflow": ("file", "critical", "lossy"),
    "layout.text_clipped": ("file", "critical", "mechanical"),
    "layout.guide_misalign": ("file", "advice", "mechanical"),
    "layout.margin_violation": ("file", "warning", "mechanical"),
    "layout.image_distorted": ("file", "warning", "mechanical"),
}


def element(
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    slot_id: str = "body",
    text: str = "Пилот окупился за четыре квартала",
    size: float = 18,
) -> RenderedElement:
    return RenderedElement(
        slot_id=slot_id,
        kind="text",
        bounds=Bounds(x=x, y=y, w=w, h=h),
        runs=[TextRun(text=text, font="Play", size_pt=size)],
    )


def deck(*elements: RenderedElement, layout_id: str = "l1") -> RenderedPresentation:
    return RenderedPresentation(
        variant="A",
        template_hash="0" * 8,
        canvas=CANVAS,
        slides=[Slide(id="s1", layout_id=layout_id, elements=list(elements))],
    )


def slot(slot_id: str, *, x: float, y: float, w: float, h: float) -> Slot:
    return Slot(id=slot_id, kind="body", origin="placeholder", bounds=Bounds(x=x, y=y, w=w, h=h))


def schema(*slots: Slot, layout_id: str = "l1") -> TemplateSchema:
    return TemplateSchema(
        source=TemplateSource(file="synthetic.pptx", hash="0" * 8),
        canvas=CANVAS,
        layouts=[Layout(id=layout_id, name="Макет", family="content", slots=list(slots))],
    )


def findings(check_id: str, context: AuditContext) -> list:
    """Выполнить одну проверку через каркас и вернуть её находки."""
    discover()
    report = run_checks(context)
    assert check_id in report.checks_run, f"{check_id} не выполнилась: {report.checks_skipped}"
    return [item for item in report.findings if item.check_id == check_id]


# --- Реестр --------------------------------------------------------------


@pytest.mark.parametrize("check_id", sorted(CHECKLIST))
def test_check_is_registered_as_the_checklist_describes(check_id: str) -> None:
    """Семь проверок вёрстки есть в реестре и классифицированы как в чек-листе."""
    discover()
    registered = REGISTRY.get(check_id)
    assert registered is not None, f"{check_id} не зарегистрирована"

    spec = registered.spec
    assert (spec.check_class, spec.severity, spec.fixability) == CHECKLIST[check_id]
    assert spec.category == "layout"


# --- layout.out_of_bounds -------------------------------------------------


def test_out_of_bounds_fires_on_element_past_the_canvas() -> None:
    """Элемент, вышедший за холст, виден не целиком — это дефект слайда."""
    found = findings(
        "layout.out_of_bounds",
        AuditContext(deck=deck(element(x=0.7, y=0.3, w=0.5, h=0.2))),
    )

    assert len(found) == 1
    assert found[0].slide_number == 1 and found[0].slot_id == "body"
    assert found[0].severity == "critical"


def test_out_of_bounds_is_silent_on_element_inside_the_canvas() -> None:
    assert findings("layout.out_of_bounds", AuditContext(deck=deck(element(x=0.1, y=0.3, w=0.5, h=0.2)))) == []


# --- layout.overlap -------------------------------------------------------


def test_overlap_fires_on_two_blocks_sharing_area() -> None:
    """Наложение блоков — дефект, который на рендере выглядит как каша."""
    found = findings(
        "layout.overlap",
        AuditContext(
            deck=deck(
                element(x=0.1, y=0.1, w=0.4, h=0.4, slot_id="title"),
                element(x=0.3, y=0.2, w=0.4, h=0.4, slot_id="body"),
            )
        ),
    )

    assert len(found) == 1
    assert "title" in found[0].message and "body" in found[0].message


def test_overlap_is_silent_on_blocks_placed_side_by_side() -> None:
    """Общая граница — не наложение: соседние слоты шаблона часто смежны."""
    found = findings(
        "layout.overlap",
        AuditContext(
            deck=deck(
                element(x=0.1, y=0.1, w=0.4, h=0.4, slot_id="left"),
                element(x=0.5, y=0.1, w=0.4, h=0.4, slot_id="right"),
            )
        ),
    )

    assert found == []


def test_overlap_ignores_intersection_below_the_tolerance() -> None:
    """Допуск нужен: доли пересчитываются из EMU, и точное касание редко.

    Пересечение в сотые доли процента площади не видно на слайде, а находка
    о нём обесценивала бы проверку.
    """
    found = findings(
        "layout.overlap",
        AuditContext(
            deck=deck(
                element(x=0.1, y=0.1, w=0.4, h=0.4, slot_id="left"),
                element(x=0.4999, y=0.1, w=0.4, h=0.4, slot_id="right"),
            )
        ),
    )

    assert found == []


# --- layout.text_overflow -------------------------------------------------


def test_text_overflow_fires_when_text_does_not_fit_its_frame() -> None:
    """Текст, не поместившийся в рамку, теряется при показе.

    Вместимость считается той же метрикой, что и в вёрстке: две разные
    метрики разошлись бы, и аудит объявлял бы дефектом то, что вёрстка
    считает нормой.
    """
    long_text = "Пилот окупился за четыре квартала, а внедрение заняло два месяца. " * 8
    found = findings(
        "layout.text_overflow",
        AuditContext(deck=deck(element(x=0.1, y=0.1, w=0.2, h=0.08, text=long_text, size=18))),
    )

    assert len(found) == 1
    assert found[0].severity == "critical" and found[0].fixability == "lossy"


def test_text_overflow_is_silent_when_text_fits() -> None:
    found = findings(
        "layout.text_overflow",
        AuditContext(deck=deck(element(x=0.1, y=0.1, w=0.7, h=0.3, text="Итоги пилота", size=18))),
    )

    assert found == []


# --- layout.text_clipped --------------------------------------------------


def test_text_clipped_fires_when_the_canvas_edge_cuts_the_text() -> None:
    """Край холста срезал часть текста — содержание потеряно, а не смещено."""
    found = findings(
        "layout.text_clipped",
        AuditContext(
            deck=deck(
                element(
                    x=0.8,
                    y=0.3,
                    w=0.35,
                    h=0.12,
                    text="Пилот окупился за четыре квартала и продлён на второй год",
                    size=18,
                )
            )
        ),
    )

    assert len(found) == 1
    assert found[0].slide_number == 1


def test_text_clipped_is_silent_when_the_visible_part_holds_the_text() -> None:
    """Рамка за краем — ещё не потеря текста.

    Это и отличает проверку от `out_of_bounds`: та говорит о рамке, эта — о
    содержании. Дублировать одну находку двумя сообщениями значило бы
    заставить пользователя дважды чинить одно и то же.
    """
    found = findings(
        "layout.text_clipped",
        AuditContext(deck=deck(element(x=0.1, y=0.3, w=0.95, h=0.3, text="Итоги", size=14))),
    )

    assert found == []


# --- layout.guide_misalign ------------------------------------------------


def test_guide_misalign_fires_on_block_off_the_reconstructed_guide() -> None:
    """Направляющие восстанавливаются по слотам макета, а не берутся из файла.

    `ppt/viewProps.xml` есть лишь в одном калибровочном шаблоне из трёх.
    Поэтому находка идёт как рекомендация и с пониженной уверенностью:
    обвинять в нарушении правила, выведенного кластеризацией, нельзя.
    """
    template = schema(
        slot("title", x=0.05, y=0.08, w=0.6, h=0.15),
        slot("body", x=0.05, y=0.30, w=0.6, h=0.40),
        slot("aside", x=0.05, y=0.75, w=0.6, h=0.15),
    )
    found = findings(
        "layout.guide_misalign",
        AuditContext(deck=deck(element(x=0.09, y=0.30, w=0.6, h=0.40)), template=template),
    )

    assert len(found) == 1
    assert found[0].severity == "advice"
    assert found[0].confidence < 1.0, "правило выведено кластеризацией, а не прочитано из файла"


def test_guide_misalign_is_silent_on_block_sitting_on_the_guide() -> None:
    template = schema(
        slot("title", x=0.05, y=0.08, w=0.6, h=0.15),
        slot("body", x=0.05, y=0.30, w=0.6, h=0.40),
        slot("aside", x=0.05, y=0.75, w=0.6, h=0.15),
    )
    found = findings(
        "layout.guide_misalign",
        AuditContext(deck=deck(element(x=0.05, y=0.30, w=0.6, h=0.40)), template=template),
    )

    assert found == []


# --- layout.margin_violation ----------------------------------------------


def test_margin_violation_fires_when_content_enters_the_edge_field() -> None:
    """Поля выводятся из самого шаблона: своих значений у нас нет.

    `designTokens.spacing` в схеме отсутствует, и брать поле из головы значило
    бы мерить чужой шаблон нашей линейкой.
    """
    template = schema(
        slot("title", x=0.06, y=0.08, w=0.6, h=0.15),
        slot("body", x=0.06, y=0.30, w=0.6, h=0.40),
        slot("aside", x=0.06, y=0.75, w=0.6, h=0.15),
    )
    found = findings(
        "layout.margin_violation",
        AuditContext(deck=deck(element(x=0.004, y=0.30, w=0.6, h=0.40)), template=template),
    )

    assert len(found) == 1
    assert found[0].severity == "warning"


def test_margin_violation_does_not_blame_the_template_for_its_own_slot() -> None:
    """Содержимое лежит там, где его поместил шаблон, — это не наше нарушение.

    Иначе система обвиняла бы пользователя в решениях автора шаблона: в двух
    калибровочных шаблонах есть слоты вплотную к краю холста.
    """
    template = schema(
        slot("title", x=0.06, y=0.08, w=0.6, h=0.15),
        slot("body", x=0.002, y=0.30, w=0.6, h=0.40),
        slot("aside", x=0.06, y=0.75, w=0.6, h=0.15),
    )
    found = findings(
        "layout.margin_violation",
        AuditContext(deck=deck(element(x=0.002, y=0.30, w=0.6, h=0.40)), template=template),
    )

    assert found == []


# --- layout.image_distorted -----------------------------------------------


def picture_deck(path: Path, image: Path, *, width_emu: int, height_emu: int) -> Path:
    """Собрать `.pptx` с картинкой заданного размера на слайде."""
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(CANVAS.width_emu), Emu(CANVAS.height_emu)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(image), Emu(1000000), Emu(1000000), Emu(width_emu), Emu(height_emu))
    presentation.save(str(path))
    return path


@pytest.fixture(scope="module")
def wide_image(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Картинка 800 × 400: соотношение сторон 2:1."""
    path = tmp_path_factory.mktemp("images") / "wide.png"
    Image.new("RGB", (800, 400), (30, 90, 180)).save(path)
    return path


def test_image_distorted_fires_on_a_stretched_picture(tmp_path: Path, wide_image: Path) -> None:
    """Картинка 2:1, поставленная квадратом, растянута вдвое."""
    exported = picture_deck(tmp_path / "stretched.pptx", wide_image, width_emu=4000000, height_emu=4000000)
    found = findings("layout.image_distorted", AuditContext(pptx_path=exported))

    assert len(found) == 1
    assert found[0].severity == "warning"
    assert found[0].slide_number == 1


def test_image_distorted_is_silent_on_a_picture_in_its_own_proportions(
    tmp_path: Path, wide_image: Path
) -> None:
    exported = picture_deck(tmp_path / "clean.pptx", wide_image, width_emu=4000000, height_emu=2000000)

    assert findings("layout.image_distorted", AuditContext(pptx_path=exported)) == []


# --- Вторая половина критерия на живых данных -----------------------------


@pytest.mark.parametrize(
    "name",
    [
        "VK Tech шаблон.pptx",
        "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
        "Шаблон презентации VK Education.pptx",
    ],
)
def test_clean_deck_from_a_real_template_raises_no_layout_findings(name: str) -> None:
    """На честно свёрстанной колоде проверки вёрстки молчат.

    Это вторая половина критерия приёмки, и проверять её надо на живых
    шаблонах: синтетический чистый слайд мы построили сами и знаем, что он
    чист. Здесь колода собирается тем же кодом, что пойдёт на защиту.
    """
    template_file = CALIBRATION / name
    if not template_file.exists():
        pytest.skip(f"нет калибровочного шаблона {name}")

    template = parse_template(template_file)
    structure = PresentationStructure(
        slides=[
            StructureSlide(id="s1", role="title", headline="Итоги пилота"),
            StructureSlide(
                id="s2",
                role="data",
                headline="Что изменилось",
                body=SlideBody(items=["Сборка колоды втрое быстрее", "Правки переживают сохранение"]),
            ),
            StructureSlide(
                id="s3",
                role="process",
                headline="Как шли",
                body=SlideBody(items=["Разбор", "Вёрстка", "Аудит"]),
            ),
        ]
    )

    discover()
    for variant in compose_variants(structure, template):
        report = run_checks(AuditContext(deck=variant, template=template), classes=("file",))
        layout_findings = [item for item in report.findings if item.category == "layout"]
        assert layout_findings == [], (
            f"вариант {variant.variant}: "
            + "; ".join(f"{item.check_id} — {item.message}" for item in layout_findings)
        )
