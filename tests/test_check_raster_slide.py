"""T-10: проверка `integrity.raster_slide`.

Критерий приёмки задачи: на слайде-картинке срабатывает, на нормальном — нет.

Проверка отвечает прямому требованию ТЗ (FR-38): слайд, выгруженный единым
растровым изображением, не засчитывается. Класс `file`, критичность
`critical`, исправимость `none` — растровый слайд нельзя починить правкой,
его можно только собрать заново.

Главная опасность здесь — не пропуск, а ложное срабатывание. В одном из
калибровочных шаблонов 62% макетов используют изображение как фон, и слайд
с картинкой во весь холст плюс текстом поверх — совершенно нормальный слайд.
Проверка, объявляющая его дефектом, обесценит себя за один прогон.
"""

from __future__ import annotations

import os
from collections.abc import Iterator
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from dpd.audit.checks import check_raster_slide
from dpd.export import export_pptx
from dpd.layout import compose
from dpd.models import PresentationStructure, SlideBody, StructureSlide
from dpd.parsing import parse_template

CALIBRATION = Path(__file__).resolve().parents[1] / "assets" / "templates" / "calibration"
TEMPLATE = "VK Tech шаблон.pptx"

WIDTH_EMU, HEIGHT_EMU = 12192000, 6858000


@pytest.fixture(scope="module")
def out_dir() -> Iterator[Path]:
    base = Path(os.environ.get("DPD_OUTPUT_DIR", "D:/dpd-out"))
    path = base / "tests" / "raster"
    path.mkdir(parents=True, exist_ok=True)
    yield path


@pytest.fixture(scope="module")
def picture(out_dir: Path) -> Path:
    """Изображение размером со слайд — из него собираются растровые слайды."""
    path = out_dir / "slide.png"
    Image.new("RGB", (1280, 720), (30, 90, 180)).save(path)
    return path


def build_deck(path: Path, picture: Path, *, with_text: bool, full_bleed: bool = True) -> Path:
    """Собрать `.pptx` из картинки, при необходимости добавив текст поверх."""
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(WIDTH_EMU), Emu(HEIGHT_EMU)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # пустой макет

    width = WIDTH_EMU if full_bleed else WIDTH_EMU // 3
    height = HEIGHT_EMU if full_bleed else HEIGHT_EMU // 3
    slide.shapes.add_picture(str(picture), Emu(0), Emu(0), Emu(width), Emu(height))

    if with_text:
        box = slide.shapes.add_textbox(Emu(600000), Emu(600000), Emu(4000000), Emu(1000000))
        box.text_frame.text = "Пилот окупился за четыре квартала"

    presentation.save(str(path))
    return path


def test_raster_slide_is_detected(out_dir: Path, picture: Path) -> None:
    """Слайд-картинка без текста — ровно то, что ТЗ не засчитывает."""
    deck = build_deck(out_dir / "raster.pptx", picture, with_text=False)
    findings = check_raster_slide(deck)
    assert len(findings) == 1
    assert findings[0].check_id == "integrity.raster_slide"
    assert findings[0].severity == "critical"
    assert findings[0].fixability == "none"


def test_picture_with_text_over_it_is_not_a_finding(out_dir: Path, picture: Path) -> None:
    """Фон-изображение с текстом поверх — норма, а не дефект.

    В VK Tech фоном служит изображение в 62% макетов. Срабатывание здесь
    означало бы, что проверку придётся отключить.
    """
    deck = build_deck(out_dir / "picture-with-text.pptx", picture, with_text=True)
    assert check_raster_slide(deck) == []


def test_small_picture_without_text_is_not_a_finding(out_dir: Path, picture: Path) -> None:
    """Слайд с маленькой картинкой не выгружен растром — он просто пуст.

    Это находка `integrity.empty_slide` (T-31), а не наша: смешивать их
    нельзя, потому что исправимость у них разная.
    """
    deck = build_deck(out_dir / "small-picture.pptx", picture, with_text=False, full_bleed=False)
    assert check_raster_slide(deck) == []


def test_finding_points_at_the_slide(out_dir: Path, picture: Path) -> None:
    deck = build_deck(out_dir / "raster-pointer.pptx", picture, with_text=False)
    assert check_raster_slide(deck)[0].slide_number == 1


def test_our_own_export_is_clean(out_dir: Path) -> None:
    """Главное: колода, собранная нашим пайплайном, проверку проходит."""
    template_path = CALIBRATION / TEMPLATE
    if not template_path.exists():
        pytest.skip(f"нет калибровочного шаблона {TEMPLATE}")

    structure = PresentationStructure(
        slides=[
            StructureSlide(id="s1", role="title", headline="Программа наставничества"),
            StructureSlide(
                id="s2",
                role="data",
                headline="Пилот окупился за четыре квартала",
                body=SlideBody(kind="bullets", items=["Экономия 9,4 млн", "Затраты 2,8 млн"]),
            ),
        ]
    )
    template = parse_template(template_path)
    exported = export_pptx(
        compose(structure, template), template, template_path, out_dir / "native.pptx"
    )
    assert check_raster_slide(exported) == []


def test_missing_file_raises_a_clear_error(out_dir: Path) -> None:
    with pytest.raises(FileNotFoundError):
        check_raster_slide(out_dir / "нет-такого.pptx")
