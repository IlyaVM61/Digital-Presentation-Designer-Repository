"""T-23: генерация таблиц по токенам шаблона.

Критерий приёмки задачи: таблица из контент-пакета собирается нативным
объектом, укладывается в 7×5.

Ограничение 7×5 — из чек-листа аудита: таблица крупнее перестаёт читаться
со слайда. Лишние строки не выбрасываются молча, сокращение фиксируется
компенсацией, как и переполнение текста.

Оформление таблицы синтезируется из дизайн-токенов: в 138 слайдах-примерах
трёх калибровочных шаблонов всего четыре таблицы, и копировать оформление
не из чего.
"""

from __future__ import annotations

import os
from collections.abc import Iterator
from pathlib import Path

import pytest
from pptx import Presentation

from dpd.export import export_pptx
from dpd.layout import compose
from dpd.models import PresentationStructure, StructureSlide, TableSpec, Visualization
from dpd.parsing import parse_template

CALIBRATION = Path(__file__).resolve().parents[1] / "assets" / "templates" / "calibration"
TEMPLATE = "VK Tech шаблон.pptx"

# Таблица из `assets/content-pack/content.md`.
HEADERS = ["Квартал", "Активных пар"]
ROWS = [
    ["IV кв. 2025", "18"],
    ["I кв. 2026", "34"],
    ["II кв. 2026", "61"],
    ["III кв. 2026", "96"],
]


@pytest.fixture
def template_path() -> Path:
    path = CALIBRATION / TEMPLATE
    if not path.exists():
        pytest.skip(f"нет калибровочного шаблона {TEMPLATE}")
    return path


@pytest.fixture
def out_dir(tmp_path_factory) -> Iterator[Path]:
    base = Path(os.environ.get("DPD_OUTPUT_DIR", "D:/dpd-out")) / "tests" / "tables"
    base.mkdir(parents=True, exist_ok=True)
    yield base


def deck(headers=None, rows=None) -> PresentationStructure:
    return PresentationStructure(
        slides=[
            StructureSlide(
                id="s1",
                role="data",
                headline="Рост числа пар",
                visualization=Visualization(
                    kind="table",
                    table=TableSpec(headers=headers or HEADERS, rows=rows or ROWS),
                ),
            )
        ]
    )


def test_table_element_is_produced(template_path: Path) -> None:
    rendered = compose(deck(), parse_template(template_path))
    tables = [e for s in rendered.slides for e in s.elements if e.kind == "table"]
    assert len(tables) == 1
    assert tables[0].table.headers == HEADERS


def test_table_keeps_the_content(template_path: Path) -> None:
    table = next(
        e for s in compose(deck(), parse_template(template_path)).slides
        for e in s.elements if e.kind == "table"
    )
    assert table.table.rows == ROWS


def test_oversized_table_is_trimmed_to_the_limit(template_path: Path) -> None:
    """Критерий приёмки: таблица укладывается в 7×5."""
    rows = [[f"строка {n}", str(n), "x", "y", "z", "w", "v", "u"] for n in range(20)]
    headers = [f"колонка {n}" for n in range(8)]
    rendered = compose(deck(headers, rows), parse_template(template_path))
    table = next(e for s in rendered.slides for e in s.elements if e.kind == "table")
    assert len(table.table.headers) <= 5
    assert len(table.table.rows) <= 6, "вместе с шапкой строк не больше семи"
    for row in table.table.rows:
        assert len(row) == len(table.table.headers)


def test_trimming_is_recorded(template_path: Path) -> None:
    """Потеря данных не бывает молчаливой."""
    rows = [[f"строка {n}", str(n)] for n in range(20)]
    slide = compose(deck(HEADERS, rows), parse_template(template_path)).slides[0]
    kinds = {compensation.kind for compensation in slide.applied_compensations}
    assert "tableTrim" in kinds


def test_small_table_is_untouched(template_path: Path) -> None:
    slide = compose(deck(), parse_template(template_path)).slides[0]
    assert not any(c.kind == "tableTrim" for c in slide.applied_compensations)


def test_table_is_styled_from_tokens(template_path: Path) -> None:
    """Оформление синтезируется из токенов: копировать его не из чего."""
    schema = parse_template(template_path)
    table = next(
        e for s in compose(deck(), schema).slides for e in s.elements if e.kind == "table"
    )
    fonts = {token.family for token in schema.design_tokens.fonts}
    assert table.table.font in fonts


def test_exported_table_is_native(template_path: Path, out_dir: Path) -> None:
    """Критерий приёмки: таблица — нативный объект, а не картинка и не текст."""
    schema = parse_template(template_path)
    exported = export_pptx(
        compose(deck(), schema), schema, template_path, out_dir / "table.pptx"
    )
    slide = Presentation(str(exported)).slides[0]
    tables = [shape for shape in slide.shapes if shape.has_table]
    assert tables, "нативной таблицы на слайде нет"

    grid = tables[0].table
    assert len(grid.columns) == len(HEADERS)
    assert len(grid.rows) == len(ROWS) + 1
    assert grid.cell(0, 0).text == HEADERS[0]
    assert grid.cell(1, 1).text == ROWS[0][1]


def test_generation_is_deterministic(template_path: Path) -> None:
    schema = parse_template(template_path)
    first = compose(deck(), schema).model_dump_json(by_alias=True)
    second = compose(deck(), schema).model_dump_json(by_alias=True)
    assert first == second
