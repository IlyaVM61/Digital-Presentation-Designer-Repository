"""T-08: экспорт в PPTX нативными объектами — главный технический риск проекта.

Критерий приёмки задачи: файл открывается в PowerPoint, текст выделяется и
редактируется, слайд не является изображением.

«Открывается в PowerPoint» проверить автоматически нельзя, поэтому проверяются
следствия, которые проверить можно: пакет OOXML читается сторонним разбором,
текст лежит в текстовых рамках, а не в картинке, и его действительно можно
изменить и сохранить. Требование ТЗ прямое: слайд, выгруженный растром, не
засчитывается.

Результаты пишутся в `DPD_OUTPUT_DIR` (по умолчанию `D:/dpd-out`): копия
шаблона весит десятки мегабайт, а на системном диске свободно меньше 5 ГБ.
"""

from __future__ import annotations

import os
import shutil
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from dpd.export import export_pptx
from dpd.layout import compose
from dpd.models import PresentationStructure, SlideBody, StructureSlide
from dpd.parsing import parse_template

CALIBRATION = Path(__file__).resolve().parents[1] / "assets" / "templates" / "calibration"
TEMPLATE = "VK Tech шаблон.pptx"

HEADLINE = "Пилот окупился за четыре квартала"
ITEMS = ["Экономия 9,4 млн рублей", "Затраты 2,8 млн рублей"]

DRAWINGML = "http://schemas.openxmlformats.org/drawingml/2006/main"
PML = "http://schemas.openxmlformats.org/presentationml/2006/main"


@pytest.fixture
def out_dir(tmp_path_factory) -> Path:
    """Каталог результатов: на диске с местом, а не в системном temp."""
    base = Path(os.environ.get("DPD_OUTPUT_DIR", "D:/dpd-out"))
    base.mkdir(parents=True, exist_ok=True)
    path = Path(str(base / "tests")) / tmp_path_factory.mktemp("export", numbered=True).name
    path.mkdir(parents=True, exist_ok=True)
    yield path
    shutil.rmtree(path, ignore_errors=True)


@pytest.fixture
def template_path() -> Path:
    path = CALIBRATION / TEMPLATE
    if not path.exists():
        pytest.skip(f"нет калибровочного шаблона {TEMPLATE}")
    return path


@pytest.fixture
def structure() -> PresentationStructure:
    return PresentationStructure(
        slides=[
            StructureSlide(id="s1", role="title", headline="Программа наставничества"),
            StructureSlide(
                id="s2",
                role="data",
                headline=HEADLINE,
                body=SlideBody(kind="bullets", items=ITEMS),
            ),
        ]
    )


@pytest.fixture
def exported(structure, template_path, out_dir) -> Path:
    template = parse_template(template_path)
    rendered = compose(structure, template)
    return export_pptx(rendered, template, template_path, out_dir / "deck.pptx")


def test_file_is_created(exported: Path) -> None:
    assert exported.is_file()
    assert exported.stat().st_size > 0


def test_package_is_a_valid_ooxml_zip(exported: Path) -> None:
    """«Файл открывается» начинается с того, что пакет не повреждён."""
    with zipfile.ZipFile(exported) as package:
        assert package.testzip() is None
        assert "ppt/presentation.xml" in package.namelist()


def test_example_slides_of_the_template_are_removed(exported: Path, structure) -> None:
    """Шаблон несёт десятки слайдов-примеров — в колоде остаются только наши."""
    assert len(Presentation(str(exported)).slides) == len(structure.slides)


def test_text_is_present_and_lives_in_text_frames(exported: Path) -> None:
    prs = Presentation(str(exported))
    texts = [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    joined = "\n".join(texts)
    assert HEADLINE in joined
    for item in ITEMS:
        assert item in joined


def test_no_slide_is_a_picture(exported: Path) -> None:
    """Прямое требование ТЗ: слайд, выгруженный растром, не засчитывается."""
    prs = Presentation(str(exported))
    for slide in prs.slides:
        pictures = slide.shapes._spTree.findall(f"{{{PML}}}pic")
        assert not pictures, "на слайде есть изображение"


def test_text_is_native_drawingml(exported: Path) -> None:
    """Текст лежит в `a:t` — как текст, а не как пиксели."""
    with zipfile.ZipFile(exported) as package:
        slide_xml = package.read("ppt/slides/slide2.xml").decode("utf-8")
    assert "<a:t>" in slide_xml or "<a:t " in slide_xml
    assert HEADLINE in slide_xml


def test_text_is_editable(exported: Path, out_dir: Path) -> None:
    """Сильнейшая из доступных проверок: текст меняется и переживает сохранение."""
    prs = Presentation(str(exported))
    target = next(
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame and HEADLINE in shape.text_frame.text
    )
    target.text_frame.paragraphs[0].runs[0].text = "Отредактировано"
    edited = out_dir / "edited.pptx"
    prs.save(str(edited))

    reopened = "\n".join(
        shape.text_frame.text
        for slide in Presentation(str(edited)).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "Отредактировано" in reopened
    assert HEADLINE not in reopened


def test_slides_keep_the_template_layouts(exported: Path, template_path: Path) -> None:
    """Слайд собран на макете шаблона — иначе оформление взялось бы ниоткуда."""
    template_layout_names = {
        layout.name for master in Presentation(str(template_path)).slide_masters
        for layout in master.slide_layouts
    }
    for slide in Presentation(str(exported)).slides:
        assert slide.slide_layout.name in template_layout_names


def test_fractions_are_converted_to_emu(exported: Path, structure, template_path) -> None:
    """Доли пересчитываются в EMU под холст шаблона — и только здесь."""
    rendered = compose(structure, parse_template(template_path))
    element = next(e for s in rendered.slides for e in s.elements)
    canvas = rendered.canvas

    prs = Presentation(str(exported))
    shapes = [sh for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]
    assert any(
        abs(shape.left - element.bounds.x * canvas.width_emu) <= 1 for shape in shapes
    ), "ни одна рамка не встала на пересчитанную из долей позицию"


ALL_TEMPLATES = [
    "VK Tech шаблон.pptx",
    "VK_WorkSpace_Клиентская_конференция_Шаблон_03.pptx",
    "Шаблон презентации VK Education.pptx",
]


@pytest.mark.parametrize("name", ALL_TEMPLATES)
def test_native_export_works_on_every_calibration_template(
    name: str, structure, out_dir: Path
) -> None:
    """Контрольная проверка риска: нативный экспорт на каждом из трёх шаблонов.

    Холсты у них разные, число мастеров разное, разметка разного качества.
    Если нативность держится только на одном файле, она не держится.
    """
    path = CALIBRATION / name
    if not path.exists():
        pytest.skip(f"нет калибровочного шаблона {name}")

    template = parse_template(path)
    exported = export_pptx(
        compose(structure, template), template, path, out_dir / f"{path.stem}.pptx"
    )

    prs = Presentation(str(exported))
    assert len(prs.slides) == len(structure.slides)

    texts = "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert HEADLINE in texts, "заголовок не попал в текстовую рамку"

    for slide in prs.slides:
        assert not slide.shapes._spTree.findall(f"{{{PML}}}pic"), "слайд содержит изображение"
