"""T-69: тело в рамке, которую создаёт вёрстка, — маркеры шаблона, кегль, высота.

Прогон T-66 показал на корпоративном шаблоне тело в три строки мелким
кеглем у верхнего края: слайд заполнен на десятую часть, маркеров нет, хотя
шаблон размечает списки маркерами, а место под текст сужено подписями
колонтитула внизу до середины ширины. Причины общие, а не свойства шаблона:

1. Маркер списка не извлекался из шаблона, а рамка, созданная вёрсткой,
   своего оформления абзацев не имеет.
2. Кегль тела в созданной рамке брался со ступени шкалы и не рос, даже когда
   места втрое больше, чем нужно.
3. Мелкий элемент у нижнего края сужал место по всей высоте, хотя
   укоротить место снизу дешевле; разделительная линия во всю ширину
   границей места не считалась — у неё нулевая площадь.

Шаблон синтетический, из стандартного шаблона python-pptx: тест проверяет
правило, а не подгонку под конкретный файл.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.util import Pt

from dpd.export import export_pptx
from dpd.layout.composer import compose
from dpd.layout.filling import SPREAD_FILL
from dpd.layout.overflow import room_height, text_height
from dpd.layout.styling import style_for
from dpd.models import PresentationStructure, SlideBody, StructureSlide, TemplateSchema
from dpd.parsing import parse_template
from dpd.parsing.slots import DERIVED_MARGIN, GAP

SCALE = [10, 11, 12, 13, 14, 16, 18, 20, 24, 28, 32, 36, 40, 44]
"""Кегли образца: шкала, по которой телу есть куда расти."""

MARKER, MARKER_COLOUR, MARKER_INDENT = "▪", "C00000", 228600

FOOTER_TOP = 0.86
LINE_TOP = 0.7

ITEMS = ["Первый пункт", "Второй пункт", "Третий пункт"]


def _layout(prs, name: str):
    return next(layout for layout in prs.slide_layouts if layout.name == name)


def _move(prs, target, shape) -> None:
    """Фигура создаётся на черновом слайде и переносится на макет."""
    target.shapes._spTree.append(shape._element)


def _scratch(prs):
    return prs.slides.add_slide(_layout(prs, "Blank"))


def _sample(prs) -> None:
    """Слайд-образец: шкала кеглей и три абзаца с маркером шаблона."""
    slide = prs.slides.add_slide(_layout(prs, "Blank"))
    frame = slide.shapes.add_textbox(0, 0, prs.slide_width // 2, prs.slide_height // 2).text_frame
    for number, size in enumerate(SCALE):
        paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = f"Образец {size}"
        run.font.size = Pt(size)
        if number < 3:
            ppr = paragraph._p.get_or_add_pPr()
            ppr.set("marL", str(MARKER_INDENT))
            ppr.set("indent", str(-MARKER_INDENT))
            colour = etree.SubElement(ppr, qn("a:buClr"))
            etree.SubElement(colour, qn("a:srgbClr"), val=MARKER_COLOUR)
            etree.SubElement(ppr, qn("a:buFont"), typeface="Wingdings")
            etree.SubElement(ppr, qn("a:buChar"), char=MARKER)


def _box(prs, scratch, *, x: float, y: float, w: float, h: float, text: str):
    shape = scratch.shapes.add_textbox(
        int(prs.slide_width * x), int(prs.slide_height * y), int(prs.slide_width * w), int(prs.slide_height * h)
    )
    shape.text_frame.text = text
    return shape


def _with_footer(path: Path) -> Path:
    """Макет из заголовка и двух мелких подписей у нижнего края."""
    prs = Presentation()
    scratch = _scratch(prs)
    target = _layout(prs, "Title Only")
    _move(prs, target, _box(prs, scratch, x=0.05, y=FOOTER_TOP, w=0.15, h=0.05, text="ПОДПИСЬ"))
    _move(prs, target, _box(prs, scratch, x=0.80, y=FOOTER_TOP, w=0.15, h=0.05, text="ЕЩЁ ПОДПИСЬ"))
    _sample(prs)
    prs.save(path)
    return path


def _with_line(path: Path) -> Path:
    """Макет из заголовка и разделительной линии во всю ширину."""
    prs = Presentation()
    scratch = _scratch(prs)
    line = scratch.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        int(prs.slide_width * 0.05),
        int(prs.slide_height * LINE_TOP),
        int(prs.slide_width * 0.95),
        int(prs.slide_height * LINE_TOP),
    )
    _move(prs, _layout(prs, "Title Only"), line)
    _sample(prs)
    prs.save(path)
    return path


def _only(schema: TemplateSchema, name: str) -> TemplateSchema:
    return schema.model_copy(update={"layouts": [layout for layout in schema.layouts if layout.name == name]})


def _deck(items: list[str]) -> PresentationStructure:
    return PresentationStructure(
        slides=[StructureSlide(id="s1", role="argument", headline="Заголовок", body=SlideBody(items=items))]
    )


@pytest.fixture(scope="module")
def footer_path(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return _with_footer(tmp_path_factory.mktemp("body") / "footer.pptx")


@pytest.fixture(scope="module")
def schema(footer_path: Path) -> TemplateSchema:
    return parse_template(footer_path, use_cache=False)


@pytest.fixture(scope="module")
def line_schema(tmp_path_factory: pytest.TempPathFactory) -> TemplateSchema:
    return parse_template(_with_line(tmp_path_factory.mktemp("body") / "line.pptx"), use_cache=False)


def _derived(schema: TemplateSchema):
    layout = next(layout for layout in schema.layouts if layout.name == "Title Only")
    return next(slot for slot in layout.slots if slot.origin == "derived")


def _body(slide):
    return next(element for element in slide.elements if element.kind == "text" and element.slot_id.startswith("body"))


def _title(slide):
    return next(element for element in slide.elements if element.slot_id.startswith("title"))


# --- Маркер шаблона ---------------------------------------------------------


def test_bullet_token_is_read_from_slide_markup(schema: TemplateSchema) -> None:
    bullet = schema.design_tokens.bullet
    assert (bullet.char, bullet.font, bullet.color) == (MARKER, "Wingdings", f"#{MARKER_COLOUR}")
    assert bullet.indent == pytest.approx(MARKER_INDENT / 9144000)


def test_bullet_token_falls_back_to_master(tmp_path: Path) -> None:
    path = tmp_path / "plain.pptx"
    Presentation().save(path)
    bullet = parse_template(path, use_cache=False).design_tokens.bullet
    assert bullet.char == "•"
    assert bullet.sources == ["master.txStyles"]


def test_list_in_constructed_frame_gets_template_bullet(schema: TemplateSchema) -> None:
    body = _body(compose(_deck(ITEMS), _only(schema, "Title Only")).slides[0])
    assert body.bullet is not None and body.bullet.char == MARKER
    assert body.frame.indent >= body.bullet.indent


def test_single_paragraph_has_no_bullet(schema: TemplateSchema) -> None:
    body = _body(compose(_deck(["Один вывод"]), _only(schema, "Title Only")).slides[0])
    assert body.bullet is None


# --- Кегль и высота ---------------------------------------------------------


def test_body_grows_but_stays_below_title(schema: TemplateSchema) -> None:
    slide = compose(_deck(ITEMS), _only(schema, "Title Only")).slides[0]
    base = style_for(_derived(schema), schema.design_tokens).size_pt
    size, title = _body(slide).runs[0].size_pt, _title(slide).runs[0].size_pt
    assert size > base
    assert size in schema.design_tokens.type_scale.values
    assert size * 1.4 <= title


def test_sparse_list_is_spread_over_the_room(schema: TemplateSchema) -> None:
    body = _body(compose(_deck(ITEMS), _only(schema, "Title Only")).slides[0])
    paragraphs, size = [run.text for run in body.runs], body.runs[0].size_pt
    height = text_height(paragraphs, body.bounds, schema.canvas, size, body.frame)
    room = room_height(body.bounds, schema.canvas, body.frame)
    assert body.frame.space_before_pt > 0
    assert height <= room
    assert body.anchor == ("middle" if height < SPREAD_FILL * room else "top")


def test_list_too_short_to_spread_is_centred(schema: TemplateSchema) -> None:
    """Отбивка ограничена: два пункта и с ней малы для места — блок встаёт по середине."""
    body = _body(compose(_deck(ITEMS[:2]), _only(schema, "Title Only")).slides[0])
    assert body.anchor == "middle"


def test_placeholder_body_keeps_template_size(schema: TemplateSchema) -> None:
    only = _only(schema, "Title and Content")
    slot = next(slot for slot in only.layouts[0].slots if slot.kind == "body")
    body = _body(compose(_deck(ITEMS), only).slides[0])
    assert body.runs[0].size_pt == style_for(slot, schema.design_tokens).size_pt
    assert body.bullet is None


# --- Место под текст ---------------------------------------------------------


def test_small_footer_trims_room_instead_of_narrowing(schema: TemplateSchema) -> None:
    bounds = _derived(schema).bounds
    assert bounds.x == pytest.approx(DERIVED_MARGIN)
    assert bounds.x + bounds.w == pytest.approx(1 - DERIVED_MARGIN)
    assert bounds.y + bounds.h <= FOOTER_TOP - GAP + 1e-6


def test_full_width_line_bounds_the_room(line_schema: TemplateSchema) -> None:
    bounds = _derived(line_schema).bounds
    assert bounds.y + bounds.h <= LINE_TOP - GAP + 1e-6


# --- Выгрузка ---------------------------------------------------------------


def test_export_writes_bullets_spacing_and_anchor(schema: TemplateSchema, footer_path: Path, tmp_path: Path) -> None:
    only = _only(schema, "Title Only")
    items = ITEMS[:2]
    out = export_pptx(compose(_deck(items), only), only, footer_path, tmp_path / "deck.pptx")
    slide = Presentation(out).slides[0]
    box = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.startswith(items[0]))
    paragraphs = box.text_frame._txBody.xpath("./a:p")
    assert len(paragraphs) == len(items)
    for paragraph in paragraphs:
        ppr = paragraph.find(qn("a:pPr"))
        assert ppr.find(qn("a:buChar")).get("char") == MARKER
        assert ppr.find(qn("a:buFont")).get("typeface") == "Wingdings"
        assert ppr.find(qn("a:buClr"))[0].get("val") == MARKER_COLOUR
        assert int(ppr.get("marL")) > 0 and int(ppr.get("indent")) == -int(ppr.get("marL"))
        assert ppr.find(qn("a:spcBef")) is not None
    assert box.text_frame._txBody.find(qn("a:bodyPr")).get("anchor") == "ctr"
