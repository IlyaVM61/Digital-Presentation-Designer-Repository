"""T-11: разрешение цепочки наследования свойств текста.

Критерий приёмки задачи: на VK Tech корректно разрешаются 59% прогонов без
явного кегля, `resolvedFrom` заполнен.

Цепочка из `docs/04-architecture/template-schema.md`:

    прогон → абзац → слот макета (lstStyle) → мастер (txStyles)
      → тема (fontScheme) → значение по умолчанию приложения

Отсутствие явного значения — норма, а не порча входных данных: в
калибровочных шаблонах от 25% до 59% прогонов не несут явного кегля. Парсер,
требующий явного значения, на реальном корпоративном шаблоне откажет —
именно это происходит с presenton (issue #794).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from dpd.models import TextStyle
from dpd.parsing import StyleResolver

CALIBRATION = Path(__file__).resolve().parents[1] / "assets" / "templates" / "calibration"
VK_TECH = "VK Tech шаблон.pptx"

# Измерено на фазе 2: 1056 из 1800 прогонов VK Tech без явного кегля.
# Наш обход даёт близкое, но не тождественное число прогонов, поэтому
# проверяется доля, а не абсолют.
EXPECTED_INHERITED_SHARE = (0.55, 0.65)


def open_vk_tech():
    path = CALIBRATION / VK_TECH
    if not path.exists():
        pytest.skip(f"нет калибровочного шаблона {VK_TECH}")
    return Presentation(str(path))


@pytest.fixture(scope="module")
def resolved() -> list[TextStyle]:
    """Разрешённые стили всех текстовых прогонов слайдов VK Tech."""
    presentation = open_vk_tech()
    resolver = StyleResolver(presentation)
    return [
        resolver.resolve_run(run, shape, slide, paragraph)
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]


def test_every_run_gets_a_size(resolved: list[TextStyle]) -> None:
    """Кегль находится всегда — это и есть смысл разрешения цепочки."""
    assert resolved
    missing = [style for style in resolved if style.size_pt is None]
    assert not missing, f"{len(missing)} прогонов остались без кегля"


def test_every_run_records_where_the_value_came_from(resolved: list[TextStyle]) -> None:
    """`resolvedFrom` нужен для честности отчётов: «кегль из мастера» и
    «кегль задан явно» — утверждения разной достоверности."""
    assert all(style.resolved_from for style in resolved)


def test_inherited_share_matches_the_measurement(resolved: list[TextStyle]) -> None:
    """Доля унаследованных кеглей воспроизводит замер фазы 2."""
    inherited = [style for style in resolved if style.resolved_from != "run"]
    share = len(inherited) / len(resolved)
    low, high = EXPECTED_INHERITED_SHARE
    assert low <= share <= high, f"доля наследования {share:.0%} вне ожидаемых {low:.0%}–{high:.0%}"


def test_resolution_levels_are_from_the_documented_chain(resolved: list[TextStyle]) -> None:
    allowed = {
        "run",
        "paragraph",
        "layout.lstStyle",
        "master.txStyles",
        "theme.fontScheme",
        "default",
    }
    assert {style.resolved_from for style in resolved} <= allowed


def test_explicit_run_value_wins(resolved: list[TextStyle]) -> None:
    """Первое найденное значение выигрывает: явный прогон не перебивается."""
    presentation = open_vk_tech()
    resolver = StyleResolver(presentation)
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    properties = run._r.find(f"{namespace}rPr")
                    if properties is not None and properties.get("sz"):
                        style = resolver.resolve_run(run, shape, slide, paragraph)
                        assert style.resolved_from == "run"
                        assert style.size_pt == int(properties.get("sz")) / 100
                        return
    pytest.skip("в шаблоне не нашлось прогона с явным кеглем")


def test_values_are_plausible(resolved: list[TextStyle]) -> None:
    """Разрешённый кегль должен быть похож на кегль, а не на артефакт."""
    for style in resolved:
        assert 1 <= style.size_pt <= 400


def test_fonts_are_resolved_for_most_runs(resolved: list[TextStyle]) -> None:
    """Гарнитура наследуется так же, как кегль — 60% прогонов её не несут."""
    with_font = [style for style in resolved if style.font]
    assert len(with_font) / len(resolved) > 0.9


def test_slot_styles_are_resolved_in_the_schema() -> None:
    """Слоты схемы несут разрешённый стиль: вёрстке нужен он, а не сырой XML."""
    from dpd.parsing import parse_template

    path = CALIBRATION / VK_TECH
    if not path.exists():
        pytest.skip(f"нет калибровочного шаблона {VK_TECH}")

    schema = parse_template(path)
    styled = [
        slot for layout in schema.layouts for slot in layout.slots if slot.text_style is not None
    ]
    assert styled, "ни один слот не получил разрешённый стиль"
    assert all(slot.text_style.resolved_from for slot in styled)
