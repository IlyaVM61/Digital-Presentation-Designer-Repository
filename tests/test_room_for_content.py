"""T-57: содержание не теряется из-за ложных слотов.

Прогон T-43 собрал колоду, где у 10 слайдов из 12 остался один заголовок.
Причин три, и все общие, а не свойства одного шаблона:

1. Пустая декоративная автофигура считалась слотом тела: текстовая рамка в
   python-pptx есть у любой автофигуры, и «рамка» ничего не отсекала. Макет
   с декором попадал в разделённые и не выбирался под контентный слайд.
2. Номер слайда и колонтитулы считались местом под содержимое: макет из
   заголовка и номера слайда становился контентным, а слот в свободной
   области ему не строился — тело было некуда положить.
3. Макет нужного типа без места под содержимое выбирался раньше пригодного
   макета соседнего типа.

Шаблон здесь синтетический, из стандартного шаблона python-pptx: тест
проверяет правило, а не подгонку под конкретный файл.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from dpd.layout.selector import select
from dpd.models import Bounds, Layout, SlideBody, Slot, StructureSlide
from dpd.parsing import parse_template

MASTER_DECOR_LEFT = 0.82
"""Левая граница декоративной полосы мастера, в долях ширины холста."""


def _place(prs, target, *, x: float, y: float, w: float, h: float, text: str | None = None, box: bool = False):
    """Добавить фигуру на макет или мастер.

    Макет и мастер в python-pptx фигуры не добавляют: фигура создаётся на
    черновом слайде и переносится в дерево фигур цели.
    """
    scratch = prs.slides[0] if len(prs.slides) else prs.slides.add_slide(_layout(prs, "Blank"))
    geometry = (int(prs.slide_width * x), int(prs.slide_height * y), int(prs.slide_width * w), int(prs.slide_height * h))
    shape = scratch.shapes.add_textbox(*geometry) if box else scratch.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, *geometry)
    if text:
        shape.text_frame.text = text
    target.shapes._spTree.append(shape._element)


def _layout(prs, name: str):
    return next(layout for layout in prs.slide_layouts if layout.name == name)


def _build(path: Path, *, hide_master_shapes: bool = False) -> Path:
    prs = Presentation()
    # Декоративная полоса мастера у правого края во всю высоту.
    _place(prs, prs.slide_master, x=MASTER_DECOR_LEFT, y=0.0, w=0.12, h=1.0)

    blank = _layout(prs, "Blank")
    _place(prs, blank, x=0.05, y=0.3, w=0.4, h=0.4, text="Место под текст")
    _place(prs, blank, x=0.05, y=0.75, w=0.4, h=0.15, box=True)
    for top in (0.1, 0.4, 0.7):
        _place(prs, blank, x=0.55, y=top, w=0.2, h=0.2)

    if hide_master_shapes:
        _layout(prs, "Title Only")._element.set("showMasterSp", "0")
    prs.save(path)
    return path


@pytest.fixture(scope="module")
def layouts(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Layout]:
    path = _build(tmp_path_factory.mktemp("room") / "synthetic.pptx")
    return {layout.name: layout for layout in parse_template(path, use_cache=False).layouts}


def _bodies(layout: Layout) -> list[Slot]:
    return [slot for slot in layout.slots if slot.kind == "body"]


# --- Пустой декор не слот -------------------------------------------------


def test_empty_decorative_shapes_are_not_slots(layouts: dict[str, Layout]) -> None:
    decor = [slot for slot in layouts["Blank"].slots if slot.origin == "shape" and slot.bounds.x > 0.5]
    assert decor == [], f"пустые фигуры приняты за слоты: {[slot.id for slot in decor]}"


def test_shapes_meant_for_text_are_still_slots(layouts: dict[str, Layout]) -> None:
    """Фигура с текстом и надпись — заявление автора о месте под текст."""
    found = sorted(round(slot.bounds.y, 2) for slot in layouts["Blank"].slots if slot.origin == "shape")
    assert found == [0.3, 0.75]


# --- Номер слайда — не место под содержимое -------------------------------


def test_slide_number_does_not_count_as_room_for_content(layouts: dict[str, Layout]) -> None:
    """Макет «заголовок и номер слайда» получает слот тела в свободной области."""
    title_only = layouts["Title Only"]
    assert [slot.origin for slot in _bodies(title_only)] == ["derived"]
    assert title_only.family == "content"


def test_derived_slot_steps_around_master_decor(layouts: dict[str, Layout]) -> None:
    """Фигуры мастера видны на макете — сконструированный слот их обходит."""
    (body,) = _bodies(layouts["Title Only"])
    assert body.bounds.x + body.bounds.w <= MASTER_DECOR_LEFT


def test_hidden_master_decor_is_not_an_obstacle(tmp_path: Path) -> None:
    path = _build(tmp_path / "hidden.pptx", hide_master_shapes=True)
    title_only = next(layout for layout in parse_template(path, use_cache=False).layouts if layout.name == "Title Only")
    (body,) = _bodies(title_only)
    assert body.bounds.x + body.bounds.w > MASTER_DECOR_LEFT


def test_layout_without_title_is_not_offered_for_content(layouts: dict[str, Layout]) -> None:
    """Заголовок слайда на макет без слота заголовка не ложится: он пропадает.

    Так в прогоне T-43 у двух слайдов варианта B не стало заголовка. Такой
    макет — последний запасной, а не контентный и не разделённый.
    """
    assert len(_bodies(layouts["Blank"])) == 2
    assert layouts["Blank"].family == "blank"


# --- Выбор макета ---------------------------------------------------------


def _slot(kind: str, slot_id: str, x: float, y: float, w: float, h: float) -> Slot:
    return Slot(id=slot_id, kind=kind, origin="placeholder", bounds=Bounds(x=x, y=y, w=w, h=h))


def test_layout_without_room_loses_to_a_suitable_neighbour_family() -> None:
    """Слайд с телом не кладётся на макет, где тело некуда положить."""
    title = _slot("title", "title", 0.05, 0.05, 0.9, 0.1)
    no_room = Layout(
        id="m/l1", name="Только заголовок", family="content",
        slots=[title, _slot("other", "other-1", 0.9, 0.9, 0.06, 0.08)],
    )
    split = Layout(
        id="m/l2", name="Две колонки", family="split",
        slots=[title, _slot("body", "body-1", 0.05, 0.2, 0.42, 0.7), _slot("body", "body-2", 0.53, 0.2, 0.42, 0.7)],
    )
    slide = StructureSlide(id="s1", role="data", headline="Итоги", body=SlideBody(items=["Пункт"]))

    chosen, decision = select(slide, [no_room, split])

    assert chosen.id == "m/l2"
    assert decision.degraded
