"""T-56: текст таблиц и диаграмм не сливается с фоном.

Аудит визуала (T-52) нашёл на калибровочных шаблонах шапку таблицы того же
цвета, что её заливка, белый текст строк на светлой заливке и подписи
диаграмм в цвет фона. Причина общая: **заливку ячеек и цвет подписей
диаграммы вёрстка не задавала.** Заливку брал стиль таблицы по умолчанию,
который python-pptx вписывает в каждую таблицу, — он красит шапку цветом
акцента темы, а строки его светлыми оттенками; текст при этом красила
вёрстка, и пару никто не сверял. Подписи диаграммы оставались без цвета
вовсе и получали его от программы просмотра.

Теперь пару «текст — то, что под ним» вёрстка задаёт целиком и сама:
шапка — брендовый цвет с текстом, читаемым на нём; строки — без заливки,
текстом того же цвета, что основной текст слайда; подписи диаграммы — тем
же цветом. Детерминированная проверка контраста видит эти пары и меряет
их, как текст.

Шаблоны синтетические: условие под конкретный файл здесь невозможно.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor

from dpd.audit import AuditContext, discover, run_checks
from dpd.export import export_pptx
from dpd.layout import charts, compose, tables
from dpd.models import (
    AxisTitles,
    Background,
    Bounds,
    Canvas,
    ChartSeries,
    ChartSpec,
    ColorToken,
    DesignTokens,
    Layout,
    PresentationStructure,
    RenderedChart,
    RenderedElement,
    RenderedPresentation,
    RenderedTable,
    Slide,
    Slot,
    StructureSlide,
    TableSpec,
    TemplateSchema,
    TemplateSource,
    TextStyle,
    Visualization,
)
from dpd.parsing import parse_template

BOUNDS = Bounds(x=0.1, y=0.3, w=0.6, h=0.5)
CANVAS = Canvas(width_emu=12192000, height_emu=6858000)
CHECK_ID = "template.contrast_low"

VK_BLUE = "#0077FF"
"""Белый на нём — 4,13:1, ниже 4,5:1; чёрный — 5,08:1."""


def tokens(*values: str, brand: str | None = None) -> DesignTokens:
    """Палитра в порядке частоты; `brand` получает роль брендового цвета."""
    colours = [
        ColorToken(role="brand.primary" if value == brand else f"accent.{n}", value=value)
        for n, value in enumerate(values)
    ]
    return DesignTokens(colors=colours)


def slot(colour: str | None) -> Slot:
    return Slot(
        id="body",
        kind="body",
        origin="placeholder",
        bounds=BOUNDS,
        text_style=TextStyle(color=colour, size_pt=14.0, resolved_from="layout"),
    )


def table_spec() -> TableSpec:
    return TableSpec(headers=["Показатель", "Значение"], rows=[["Затраты", "2,8 млн ₽"]])


def contrast(first: str, second: str) -> float:
    def luminance(colour: str) -> float:
        channels = [int(colour[i : i + 2], 16) / 255 for i in (1, 3, 5)]
        linear = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in channels]
        return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]

    lighter, darker = sorted((luminance(first), luminance(second)), reverse=True)
    return (lighter + 0.05) / (darker + 0.05)


# --- Вёрстка таблицы ------------------------------------------------------


def test_header_is_filled_with_the_brand_colour() -> None:
    table, _ = tables.build(table_spec(), slot("#FFFFFF"), tokens("#FFFFFF", "#0B3D91", brand="#0B3D91"))
    assert table.header_fill == "#0B3D91"


def test_header_text_is_never_the_colour_of_its_fill() -> None:
    """Дефект VK Education и VK WorkSpace: шапка брендовым по брендовому."""
    table, _ = tables.build(
        table_spec(), slot("#FFFFFF"), tokens("#FFFFFF", VK_BLUE, "#000000", brand=VK_BLUE)
    )
    assert table.header_color != table.header_fill
    assert contrast(table.header_color, table.header_fill) >= 4.5


def test_header_text_prefers_the_body_colour_when_it_reads() -> None:
    """Шапка тем же цветом, что строки, если он читается на заливке."""
    table, _ = tables.build(
        table_spec(), slot("#FFFFFF"), tokens("#000000", "#FFFFFF", "#0B3D91", brand="#0B3D91")
    )
    assert table.header_color == "#FFFFFF"


def test_header_text_comes_from_the_palette() -> None:
    """Чужой цвет выдаёт, что слайд собран мимо шаблона."""
    palette = ("#FFFFFF", VK_BLUE, "#0A0A0A")
    table, _ = tables.build(table_spec(), slot("#FFFFFF"), tokens(*palette, brand=VK_BLUE))
    assert table.header_color in palette


def test_header_loses_its_fill_when_nothing_in_the_palette_reads_on_it() -> None:
    """Без читаемой пары шапка остаётся без заливки, как строки."""
    table, _ = tables.build(
        table_spec(), slot("#FFFFFF"), tokens("#FFFFFF", VK_BLUE, "#3399FF", brand=VK_BLUE)
    )
    assert table.header_fill is None
    assert table.header_color == "#FFFFFF"


def test_header_without_a_brand_colour_has_no_fill() -> None:
    table, _ = tables.build(table_spec(), slot("#222222"), tokens("#222222", "#777777"))
    assert table.header_fill is None
    assert table.header_color == "#222222"


def test_rows_are_set_in_the_body_text_colour() -> None:
    """Строки без заливки: под ними фон слайда, и цвет текста — как у основного текста."""
    table, _ = tables.build(table_spec(), slot("#FFFFFF"), tokens("#FFFFFF", VK_BLUE, brand=VK_BLUE))
    assert table.body_color == "#FFFFFF"


# --- Вёрстка диаграммы ----------------------------------------------------


def chart_spec() -> ChartSpec:
    return ChartSpec(
        chart_type="bar",
        categories=["Текущие", "Требуемые"],
        series=[ChartSeries(name="Количество", points=[90, 310])],
        axis_titles=AxisTitles(category="Группа", value="Человек"),
    )


def test_chart_text_is_set_in_the_body_text_colour() -> None:
    """Дефект VK Tech и VK WorkSpace: подписи без цвета — чёрным по чёрному."""
    chart = charts.build(chart_spec(), slot("#FFFFFF"), tokens("#FFFFFF", VK_BLUE, brand=VK_BLUE))
    assert chart.text_color == "#FFFFFF"


# --- Экспорт --------------------------------------------------------------

CALIBRATION = Path(__file__).resolve().parents[1] / "assets" / "templates" / "calibration"
TEMPLATE = "VK Tech шаблон.pptx"


@pytest.fixture
def template_path() -> Path:
    path = CALIBRATION / TEMPLATE
    if not path.exists():
        pytest.skip(f"нет калибровочного шаблона {TEMPLATE}")
    return path


@pytest.fixture
def out_dir() -> Path:
    base = Path(os.environ.get("DPD_OUTPUT_DIR", "D:/dpd-out")) / "tests" / "readable"
    base.mkdir(parents=True, exist_ok=True)
    return base


def structure(visual: Visualization) -> PresentationStructure:
    return PresentationStructure(
        slides=[StructureSlide(id="s1", role="data", headline="Экономика пилота", visualization=visual)]
    )


def rgb(colour: str) -> RGBColor:
    return RGBColor.from_string(colour.lstrip("#"))


def test_exported_table_carries_its_own_fills(template_path: Path, out_dir: Path) -> None:
    """Заливку задаёт вёрстка, а не стиль таблицы по умолчанию."""
    schema = parse_template(template_path)
    rendered = compose(structure(Visualization(kind="table", table=table_spec())), schema)
    table = next(e for s in rendered.slides for e in s.elements if e.kind == "table").table
    exported = export_pptx(rendered, schema, template_path, out_dir / "table.pptx")

    grid = next(shape for shape in Presentation(str(exported)).slides[0].shapes if shape.has_table).table
    header, row = grid.cell(0, 0), grid.cell(1, 0)
    if table.header_fill:
        assert header.fill.fore_color.rgb == rgb(table.header_fill)
    assert row._tc.tcPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}noFill") is not None, (
        "у строки нет явного «без заливки» — её закрасит стиль таблицы"
    )
    assert header.text_frame.paragraphs[0].runs[0].font.color.rgb == rgb(table.header_color)
    assert row.text_frame.paragraphs[0].runs[0].font.color.rgb == rgb(table.body_color)


def test_exported_chart_text_has_a_colour(template_path: Path, out_dir: Path) -> None:
    schema = parse_template(template_path)
    rendered = compose(structure(Visualization(kind="chart", chart=chart_spec())), schema)
    chart_element = next(e for s in rendered.slides for e in s.elements if e.kind == "chart").chart
    exported = export_pptx(rendered, schema, template_path, out_dir / "chart.pptx")

    chart = next(shape for shape in Presentation(str(exported)).slides[0].shapes if shape.has_chart).chart
    colour = rgb(chart_element.text_color)
    assert chart.font.color.rgb == colour
    assert chart.category_axis.tick_labels.font.color.rgb == colour
    assert chart.value_axis.tick_labels.font.color.rgb == colour


# --- Детерминированная проверка контраста --------------------------------


def schema_on(background: str | None, computable: bool = True) -> TemplateSchema:
    return TemplateSchema(
        source=TemplateSource(file="synthetic.pptx", hash="0" * 8),
        canvas=CANVAS,
        layouts=[
            Layout(
                id="l1",
                name="Макет",
                family="content",
                background=Background(
                    kind="solid" if computable else "image", value=background, contrast_computable=computable
                ),
                slots=[slot(None)],
            )
        ],
    )


def deck_with(element: RenderedElement) -> RenderedPresentation:
    return RenderedPresentation(
        variant="A",
        template_hash="0" * 8,
        canvas=CANVAS,
        slides=[Slide(id="s1", layout_id="l1", elements=[element])],
    )


def table_element(header: str, fill: str | None, body: str) -> RenderedElement:
    return RenderedElement(
        slot_id="body",
        kind="table",
        bounds=BOUNDS,
        table=RenderedTable(
            headers=["Показатель"], rows=[["Затраты"]], size_pt=12.0,
            header_color=header, header_fill=fill, body_color=body,
        ),
    )


def chart_element(text: str) -> RenderedElement:
    return RenderedElement(
        slot_id="body",
        kind="chart",
        bounds=BOUNDS,
        chart=RenderedChart(categories=["Текущие"], size_pt=12.0, text_color=text),
    )


def contrast_findings(deck: RenderedPresentation, template: TemplateSchema) -> list:
    discover()
    report = run_checks(AuditContext(deck=deck, template=template))
    assert CHECK_ID in report.checks_run
    return [item for item in report.findings if item.check_id == CHECK_ID]


def test_header_text_on_its_own_fill_is_found() -> None:
    """Шапка видна проверке и на фоне-изображении: заливка шапки — код цвета."""
    found = contrast_findings(deck_with(table_element(VK_BLUE, VK_BLUE, "#FFFFFF")), schema_on(None, False))
    assert len(found) == 1
    assert found[0].check_class == "file"
    assert found[0].evidence["background"] == VK_BLUE


def test_table_rows_are_measured_against_the_slide_background() -> None:
    found = contrast_findings(deck_with(table_element("#000000", VK_BLUE, "#FFFFFF")), schema_on("#F5F5F5"))
    assert [item.evidence["text"] for item in found] == ["#FFFFFF"]


def test_chart_text_is_measured_against_the_slide_background() -> None:
    found = contrast_findings(deck_with(chart_element("#111111")), schema_on("#000000"))
    assert len(found) == 1
    assert found[0].evidence["text"] == "#111111"


def test_readable_table_and_chart_are_silent() -> None:
    assert contrast_findings(deck_with(table_element("#000000", VK_BLUE, "#FFFFFF")), schema_on("#000000")) == []
    assert contrast_findings(deck_with(chart_element("#FFFFFF")), schema_on("#000000")) == []
