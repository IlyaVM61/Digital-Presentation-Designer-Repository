"""Разрешение цепочки наследования свойств текста.

Цепочка из `docs/04-architecture/template-schema.md`, первое найденное
значение выигрывает:

    прогон → абзац → слот макета (lstStyle) → мастер (txStyles)
      → тема (fontScheme) → значение по умолчанию приложения

**Это обязательное требование, а не улучшение.** Отсутствие явного значения
на верхних уровнях — норма: в калибровочных шаблонах от 25% до 59% текстовых
прогонов не несут явного кегля и гарнитуры. Парсер, требующий явного
значения, обработает такой файл неверно либо откажет — именно это происходит
с presenton (issue #794).

Уровень, на котором значение найдено, сохраняется: «кегль взят из мастера» и
«кегль задан явно» — утверждения разной достоверности, и отчёт обязан их
различать.
"""

from __future__ import annotations

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from dpd.models import TextStyle

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

DEFAULT_SIZE_PT = 18.0
DEFAULT_FONT = "Calibri"

_TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
_BODY_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.OBJECT}


class StyleResolver:
    """Разрешает свойства текста по цепочке наследования.

    Создаётся на презентацию: тема читается один раз, иначе разбор полутора
    тысяч прогонов упёрся бы в повторный обход XML темы.
    """

    def __init__(self, presentation) -> None:
        self._presentation = presentation
        self._theme_fonts: dict[int, dict[str, str]] = {}
        self._theme_colours_cache: dict[int, dict[str, str]] = {}

    def resolve_run(self, run, shape, slide, paragraph) -> TextStyle:
        """Разрешить свойства текстового прогона на слайде."""
        layout = getattr(slide, "slide_layout", None)
        chain = [
            ("run", run._r.find(A + "rPr")),
            ("paragraph", _default_properties(paragraph._p.find(A + "pPr"))),
        ]
        return self._resolve(chain, shape, layout, paragraph.level or 0)

    def resolve_slot(self, placeholder, layout) -> TextStyle:
        """Разрешить свойства слота макета — их применит вёрстка."""
        chain = []
        if placeholder.has_text_frame and placeholder.text_frame.paragraphs:
            paragraph = placeholder.text_frame.paragraphs[0]
            chain.append(("paragraph", _default_properties(paragraph._p.find(A + "pPr"))))
        return self._resolve(chain, placeholder, layout, 0)

    def resolve_body_style(self, layout) -> TextStyle:
        """Стиль основного текста мастера — для слотов, которых нет в разметке.

        Сконструированному слоту неоткуда взять кегль: своей разметки у него
        нет. Наследовать его у заголовка нельзя — заголовок крупный, и текст
        в 36 пунктов не поместится в отведённую полосу и вылезет на логотип.
        """
        master = getattr(layout, "slide_master", None)
        properties = None
        if master is not None:
            styles = master.element.find(P + "txStyles")
            if styles is not None:
                properties = _level_properties(styles.find(P + "bodyStyle"), 0)

        chain = [("master.txStyles", properties)]
        size_pt, size_from = _first(chain, _size_of)
        font, font_from = _first(chain, _font_of)
        scheme = self.theme_colours(layout)
        colour, _ = _first(chain, lambda item: _colour_of(item, scheme))

        return TextStyle(
            font=font or DEFAULT_FONT,
            size_pt=size_pt or DEFAULT_SIZE_PT,
            color=colour,
            resolved_from=size_from or "default",
            font_resolved_from=font_from or "default",
        )

    def _resolve(self, levels, shape, layout, level: int) -> TextStyle:
        chain = list(levels)
        chain.append(("layout.lstStyle", self._layout_properties(shape, layout, level)))
        chain.append(("master.txStyles", self._master_properties(shape, layout, level)))

        size_pt, size_from = _first(chain, _size_of)
        font, font_from = _first(chain, _font_of)

        if font is None:
            font = self._theme_font(shape, layout)
            font_from = "theme.fontScheme" if font else None

        if size_pt is None:
            size_pt, size_from = DEFAULT_SIZE_PT, "default"
        if font is None:
            font, font_from = DEFAULT_FONT, "default"

        scheme = self.theme_colours(layout)
        colour, _ = _first(chain, lambda properties: _colour_of(properties, scheme))

        return TextStyle(
            font=font,
            size_pt=size_pt,
            color=colour,
            resolved_from=size_from,
            font_resolved_from=font_from,
        )

    def _layout_properties(self, shape, layout, level: int):
        """`lstStyle` одноимённого плейсхолдера макета."""
        if layout is None or not _is_placeholder(shape):
            return None
        idx = shape.placeholder_format.idx
        for placeholder in layout.placeholders:
            if placeholder.placeholder_format.idx == idx:
                return _level_properties(placeholder._element.find(".//" + A + "lstStyle"), level)
        return None

    def _master_properties(self, shape, layout, level: int):
        """`txStyles` мастера: titleStyle, bodyStyle или otherStyle."""
        master = getattr(layout, "slide_master", None)
        if master is None:
            return None
        styles = master.element.find(P + "txStyles")
        if styles is None:
            return None
        return _level_properties(styles.find(P + _master_style_name(shape)), level)

    def _theme_font(self, shape, layout) -> str | None:
        """Гарнитура из темы: major для заголовков, minor для остального.

        Тема — последний и наименее надёжный источник: во всех четырёх
        проверенных шаблонах она расходится с фактическим оформлением.
        """
        master = getattr(layout, "slide_master", None)
        if master is None:
            return None
        fonts = self._theme_fonts.get(id(master))
        if fonts is None:
            fonts = _read_theme_fonts(master)
            self._theme_fonts[id(master)] = fonts
        role = "major" if _master_style_name(shape) == "titleStyle" else "minor"
        return fonts.get(role)

    def theme_colours(self, layout) -> dict[str, str]:
        """Цветовая схема темы: без неё не разрешить ссылки вида `lt1`."""
        master = getattr(layout, "slide_master", None)
        if master is None:
            return {}
        colours = self._theme_colours_cache.get(id(master))
        if colours is None:
            colours = _read_theme_colours(master)
            self._theme_colours_cache[id(master)] = colours
        return colours


def _first(chain, extract):
    """Первое найденное значение и уровень, на котором оно найдено."""
    for name, properties in chain:
        if properties is None:
            continue
        value = extract(properties)
        if value is not None:
            return value, name
    return None, None


def _is_placeholder(shape) -> bool:
    try:
        return bool(shape.is_placeholder)
    except (AttributeError, ValueError):
        return False


def _master_style_name(shape) -> str:
    if not _is_placeholder(shape):
        return "otherStyle"
    placeholder_type = shape.placeholder_format.type
    if placeholder_type in _TITLE_PLACEHOLDERS:
        return "titleStyle"
    if placeholder_type in _BODY_PLACEHOLDERS:
        return "bodyStyle"
    return "otherStyle"


def _default_properties(container):
    """`defRPr` внутри `pPr` или `lvlNpPr`."""
    return None if container is None else container.find(A + "defRPr")


def _level_properties(list_style, level: int):
    """Свойства нужного уровня списка из `lstStyle` или `txStyles`.

    При отсутствии уровня берётся первый: макет может размечать только
    `lvl1pPr`, а текст на слайде сидеть на третьем уровне списка.
    """
    if list_style is None:
        return None
    for candidate in (level, 0):
        element = list_style.find(A + "lvl" + str(candidate + 1) + "pPr")
        if element is not None:
            properties = element.find(A + "defRPr")
            if properties is not None:
                return properties
    return None


def _size_of(properties) -> float | None:
    raw = properties.get("sz")
    return int(raw) / 100 if raw else None


def _font_of(properties) -> str | None:
    latin = properties.find(A + "latin")
    if latin is None:
        return None
    typeface = latin.get("typeface")
    # Ссылка вида "+mj-lt" означает «взять из темы» — это не гарнитура,
    # и принимать её за имя шрифта нельзя.
    if not typeface or typeface.startswith("+"):
        return None
    return typeface


def _colour_of(properties, scheme: dict[str, str] | None = None) -> str | None:
    """Цвет заливки текста: прямой код либо ссылка на цветовую схему темы.

    Ссылку разрешать обязательно. Тёмные шаблоны задают светлый текст именно
    так: мастер объявляет чёрный, а плейсхолдер макета ссылается на `lt1`,
    и в теме это белый. Парсер, читающий только `srgbClr`, вернул бы чёрный
    текст на чёрном фоне — и это ровно тот дефект, который не видно в данных,
    но видно на рендере.
    """
    fill = properties.find(A + "solidFill")
    if fill is None:
        return None

    srgb = fill.find(A + "srgbClr")
    value = srgb.get("val") if srgb is not None else None
    if value and len(value) == 6:
        return "#" + value.upper()

    referenced = fill.find(A + "schemeClr")
    name = referenced.get("val") if referenced is not None else None
    if name and scheme:
        # `tx1`/`bg1` — синонимы `dk1`/`lt1` в записи макетов.
        name = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}.get(name, name)
        resolved = scheme.get(name)
        if resolved:
            return "#" + resolved.upper()
    return None


def _read_theme_fonts(master) -> dict[str, str]:
    """Гарнитуры темы мастера."""
    try:
        theme = master.part.part_related_by(RT.THEME)
    except KeyError:
        return {}

    root = etree.fromstring(theme.blob)
    scheme = root.find(".//" + A + "fontScheme")
    if scheme is None:
        return {}

    fonts: dict[str, str] = {}
    for role, tag in (("major", "majorFont"), ("minor", "minorFont")):
        element = scheme.find(A + tag)
        latin = element.find(A + "latin") if element is not None else None
        typeface = latin.get("typeface") if latin is not None else None
        if typeface:
            fonts[role] = typeface
    return fonts


def _read_theme_colours(master) -> dict[str, str]:
    """Цветовая схема темы мастера.

    В отличие от гарнитур, схему приходится читать всегда: ссылки на неё —
    обычный способ задать цвет текста в макете, и без разрешения они молча
    теряются.
    """
    try:
        theme = master.part.part_related_by(RT.THEME)
    except KeyError:
        return {}

    root = etree.fromstring(theme.blob)
    scheme = root.find(".//" + A + "clrScheme")
    if scheme is None:
        return {}

    colours: dict[str, str] = {}
    for entry in scheme:
        name = entry.tag.split("}")[1]
        direct = entry.find(A + "srgbClr")
        system = entry.find(A + "sysClr")
        value = None
        if direct is not None:
            value = direct.get("val")
        elif system is not None:
            value = system.get("lastClr")
        if value and len(value) == 6:
            colours[name] = value
    return colours
