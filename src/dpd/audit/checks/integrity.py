"""Проверки целостности — раздел «Целостность» Приложения 1 ТЗ (T-31).

Шесть проверок класса `file`. Две из них — `file_broken` и `raster_slide` —
выполняются **после экспорта**: до него нечего проверять, файла ещё нет.
Подслой у них поэтому `export`, и без пути к файлу они не выполняются, о чём
отчёт говорит прямо: пустой список находок иначе неотличим от проверенного
файла, которого на самом деле нет.

Обе отвечают на вопрос, от которого зависит ценность всего прогона: открылся
ли выгруженный файл и остался ли он набором редактируемых объектов. Растровый
слайд по ТЗ не засчитывается (FR-38), а не открывающийся файл обесценивает
работу целиком.

**Две оговорки, без которых проверки обвиняли бы не по делу:**

- Титул и разделитель состоят из одного заголовка по замыслу жанра, и
  `empty_slide` их не касается. Вопрос 5 валидации контента («есть ли на
  слайде содержание, а не только заголовок») для них отвечается «нет»
  правильно.
- Подписи осей диаграммы — смысл, а не форма: единицы измерения приходят из
  структуры колоды, то есть от модели, и вёрстка их не выдумывает.
  `chart_no_labels` ловит диаграмму, до которой этот смысл не дошёл, и не
  требует осей от круговой, где их нет.

`raster_slide` написана задачей T-10 и вобрана каркасом T-27.
"""

from __future__ import annotations

from difflib import SequenceMatcher
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE

from dpd.audit.fixer import elements_of, fixes
from dpd.audit.package import read_package
from dpd.audit.registry import CheckSpec, check, param
from dpd.models import (
    PresentationStructure,
    RenderedElement,
    RenderedPresentation,
    Slide,
    StructureSlide,
    TemplateSchema,
)
from dpd.models.audit import Finding

SPARSE_FAMILIES = frozenset({"title", "section"})
"""Семейства макетов, у которых один заголовок — норма жанра."""

SPEC = CheckSpec(
    id="integrity.raster_slide",
    category="integrity",
    check_class="file",
    severity="critical",
    fixability="none",
    sublayer="export",
    title="Слайд оказался картинкой, а не редактируемыми объектами",
)

CHECK_ID = SPEC.id


@check(SPEC)
def check_raster_slide(
    pptx_path: str | Path,
    coverage_threshold: float | None = None,
) -> list[Finding]:
    """Найти слайды, выгруженные растровым изображением вместо объектов.

    Прямое требование ТЗ (FR-38): такой слайд не засчитывается.

    Срабатывает при совпадении двух условий: на слайде нет ни одного непустого
    текста и изображение занимает почти весь холст. Одного условия мало ни в
    ту, ни в другую сторону. Изображение во весь холст с текстом поверх —
    обычный слайд на фоне-картинке, а в одном из калибровочных шаблонов такой
    фон у 62% макетов; слайд же без текста и без большой картинки просто пуст,
    и это другая находка — `integrity.empty_slide` — с другой исправимостью.
    """
    if coverage_threshold is None:
        coverage_threshold = param(SPEC.id, "coverage_threshold")

    presentation, _ = read_package(pptx_path)
    if presentation is None:
        # Файл не открывается: об этом сообщает `integrity.file_broken`,
        # а искать в нём растровые слайды нечем.
        return []

    canvas_area = presentation.slide_width * presentation.slide_height

    findings: list[Finding] = []
    for number, slide in enumerate(presentation.slides, start=1):
        shapes = list(_walk(slide.shapes))
        if any(_has_text(shape) for shape in shapes):
            continue

        coverage = max((_coverage(shape, canvas_area) for shape in shapes if _is_picture(shape)), default=0.0)
        if coverage >= coverage_threshold:
            findings.append(
                SPEC.finding(
                    f"Слайд {number} выгружен изображением на {coverage:.0%} холста "
                    f"и не содержит текста: такой слайд не засчитывается по ТЗ.",
                    slide_number=number,
                    evidence={"coverage": round(coverage, 4), "threshold": coverage_threshold},
                )
            )
    return findings


def _walk(shapes):
    """Обойти фигуры, заглядывая внутрь групп: текст может лежать и там."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def _has_text(shape) -> bool:
    return bool(shape.has_text_frame and shape.text_frame.text.strip())


def _is_picture(shape) -> bool:
    return shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE)


def _coverage(shape, canvas_area: int) -> float:
    if not shape.width or not shape.height or not canvas_area:
        return 0.0
    return min(shape.width * shape.height / canvas_area, 1.0)


FILE_BROKEN = CheckSpec(
    id="integrity.file_broken",
    category="integrity",
    check_class="file",
    severity="critical",
    fixability="none",
    sublayer="export",
    title="Файл не открывается",
)

PLACEHOLDER_TEXT = CheckSpec(
    id="integrity.placeholder_text",
    category="integrity",
    check_class="file",
    severity="critical",
    fixability="semantic",
    sublayer="4b",
    title="Остался текст-заглушка",
)

EMPTY_SLIDE = CheckSpec(
    id="integrity.empty_slide",
    category="integrity",
    check_class="file",
    severity="critical",
    fixability="semantic",
    sublayer="4b",
    title="Пустой слайд или слайд с одним заголовком",
)

CHART_NO_LABELS = CheckSpec(
    id="integrity.chart_no_labels",
    category="integrity",
    check_class="file",
    severity="warning",
    fixability="mechanical",
    sublayer="4b",
    title="У диаграммы нет подписей осей или легенды",
    plain="У диаграммы не хватает подписей",
)

DUPLICATE_SLIDES = CheckSpec(
    id="integrity.duplicate_slides",
    category="integrity",
    check_class="file",
    severity="warning",
    fixability="none",
    sublayer="4a",
    title="Два слайда дублируют друг друга",
)

AXIS_FREE_CHARTS = frozenset({"pie", "doughnut"})
"""Типы диаграмм без осей: требовать у них подписи осей бессмысленно."""


@check(FILE_BROKEN)
def check_file_broken(pptx_path: str | Path) -> list[Finding]:
    """Проверить, что выгруженный файл вообще открывается.

    Эта проверка — единственная, кому вопрос «открывается ли файл»
    принадлежит. Остальные файловые проверки на нечитаемом файле молчат:
    искать в нём растровый слайд или мерить пропорции картинок нечем, а
    дублировать одну и ту же беду четырьмя сообщениями незачем.

    Исправимость `none`: сломанный файл не чинят правкой свойства — колоду
    выгружают заново.
    """
    _, error = read_package(pptx_path)
    if error is None:
        return []
    return [
        FILE_BROKEN.finding(
            f"Файл {Path(pptx_path).name} не открывается: {error.split(':')[0]}.",
            evidence={"error": error},
        )
    ]


@check(PLACEHOLDER_TEXT)
def check_placeholder_text(
    deck: RenderedPresentation,
    markers: list[str] | None = None,
) -> list[Finding]:
    """Найти оставшиеся заглушки и следы промпта.

    Маркеры лежат в конфигурации, а не в коде: их состав зависит от языка
    колоды и от того, чем пользуется генерация, и правка списка не должна
    требовать изменений в исходниках.

    Исправимость `semantic`: заглушку нельзя заменить автоматически — на её
    месте должен появиться смысл, а это перегенерация слайда.
    """
    markers = markers if markers is not None else param(PLACEHOLDER_TEXT.id, "markers")
    lowered = [str(marker).lower() for marker in markers]

    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        for element in slide.elements:
            for value in _texts(element):
                haystack = value.lower()
                for marker in lowered:
                    if marker not in haystack:
                        continue
                    findings.append(
                        PLACEHOLDER_TEXT.finding(
                            f"Слайд {number}: в тексте осталась заглушка «{marker}».",
                            slide_number=number,
                            slot_id=element.slot_id,
                            evidence={"marker": marker, "text": value[:80]},
                        )
                    )
                    break
    return findings


@check(EMPTY_SLIDE)
def check_empty_slide(deck: RenderedPresentation, template: TemplateSchema) -> list[Finding]:
    """Найти слайды без содержания.

    Это вопрос 5 валидации контента, перенесённый решением фазы 6 в
    детерминированные: непустота слотов проверяется алгоритмом точнее, чем
    моделью по картинке, и дешевле.

    Титульные слайды и разделители исключены: один заголовок — их жанр, а не
    недосмотр. Визуализация считается содержанием, хотя текстовых прогонов в
    ней нет.
    """
    titles = {
        (layout.id, slot.id)
        for layout in template.layouts
        for slot in layout.slots
        if slot.kind == "title"
    }

    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        if _family(slide) in SPARSE_FAMILIES:
            continue
        if any(_is_content(element, slide, titles) for element in slide.elements):
            continue
        findings.append(
            EMPTY_SLIDE.finding(
                f"Слайд {number} не несёт содержания — на нём только заголовок.",
                slide_number=number,
                evidence={"elements": len(slide.elements)},
            )
        )
    return findings


@check(CHART_NO_LABELS)
def check_chart_no_labels(deck: RenderedPresentation) -> list[Finding]:
    """Найти диаграммы без подписей осей и легенды.

    Правило наше: эталона в шаблонах нет — 138 слайдов-примеров, ноль
    диаграмм (QR-4). Обоснование простое: диаграмма без единиц измерения не
    сообщает ничего, а несколько рядов без легенды неразличимы.

    У круговой диаграммы осей нет, и подписи осей у неё не требуются; легенда
    из одного пункта не требуется тоже — она занимает место и ничего не
    объясняет, поэтому вёрстка её и не ставит.
    """
    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        for element in slide.elements:
            chart = element.chart
            if chart is None:
                continue

            missing: list[str] = []
            needs_axes = chart.chart_type not in AXIS_FREE_CHARTS
            if needs_axes and not (chart.axis_titles.category and chart.axis_titles.value):
                missing.append("оси")
            if len(chart.series) > 1 and not chart.has_legend:
                missing.append("легенда")
            if not missing:
                continue

            findings.append(
                CHART_NO_LABELS.finding(
                    f"Слайд {number}: у диаграммы не подписаны {' и '.join(missing)}.",
                    slide_number=number,
                    slot_id=element.slot_id,
                    evidence={"missing": missing, "chartType": chart.chart_type},
                )
            )
    return findings


@check(DUPLICATE_SLIDES)
def check_duplicate_slides(
    structure: PresentationStructure,
    similarity: float | None = None,
) -> list[Finding]:
    """Найти слайды, дублирующие друг друга по тексту.

    Работает по замыслу колоды, а не по вёрстке: дубль — свойство содержания,
    и обнаружить его дешевле до вёрстки, чем трижды находить в каждом из
    вариантов.

    Исправимость `none`: слить два слайда или оставить оба — решение
    человека. Система обязана показать факт, а не выбрать за него.
    """
    threshold = float(similarity) if similarity is not None else float(param(DUPLICATE_SLIDES.id, "similarity"))

    slides = structure.slides
    findings: list[Finding] = []
    for first in range(len(slides)):
        for second in range(first + 1, len(slides)):
            ratio = SequenceMatcher(None, _slide_text(slides[first]), _slide_text(slides[second])).ratio()
            if ratio <= threshold:
                continue
            findings.append(
                DUPLICATE_SLIDES.finding(
                    f"Слайды {first + 1} и {second + 1} совпадают по тексту на {ratio:.0%}.",
                    slide_number=second + 1,
                    evidence={"similarity": round(ratio, 3), "with": first + 1},
                )
            )
    return findings


def _texts(element: RenderedElement) -> list[str]:
    """Весь текст элемента: прогоны, ячейки таблицы, подписи диаграммы."""
    values = [run.text for run in element.runs]
    if element.table:
        values.extend(element.table.headers)
        values.extend(cell for row in element.table.rows for cell in row)
    if element.chart:
        values.extend(element.chart.categories)
        values.extend(series.name for series in element.chart.series)
        values.extend(
            value for value in (element.chart.axis_titles.category, element.chart.axis_titles.value) if value
        )
    return [value for value in values if value]


def _family(slide: Slide) -> str:
    return slide.layout_decision.chosen_family if slide.layout_decision else ""


def _is_content(element: RenderedElement, slide: Slide, titles: set[tuple[str, str]]) -> bool:
    """Несёт ли элемент содержание слайда, а не только его заголовок."""
    if (slide.layout_id, element.slot_id) in titles:
        return False
    if element.table is not None or element.chart is not None:
        return True
    return any(run.text.strip() for run in element.runs)


def _slide_text(slide: StructureSlide) -> str:
    """Текст слайда замысла — то, что зритель прочтёт."""
    parts = [slide.headline, slide.key_message or ""]
    if slide.body:
        parts.extend(slide.body.items)
    return " ".join(part.strip() for part in parts if part).lower()


# --- Исправления (T-33) ----------------------------------------------------


@fixes(CHART_NO_LABELS)
def fix_chart_no_labels(deck, finding, template) -> bool:
    """Включить легенду. Подписи осей фиксер не выдумывает.

    Легенда — форма: рядов больше одного, значит их надо назвать, и назвать
    их нечем, кроме имён самих рядов. Единицы измерения — смысл: выдумать
    их нельзя, а подписать ось словом «значение» и отчитаться об исправлении
    было бы обманом. Поэтому находка с недостающими осями остаётся открытой.
    """
    if "легенда" not in finding.evidence.get("missing", []):
        return False

    changed = False
    for element in elements_of(deck, finding):
        if element.chart is not None and not element.chart.has_legend:
            element.chart.has_legend, changed = True, True
    return changed
