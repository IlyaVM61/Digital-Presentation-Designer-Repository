"""Проверки вёрстки — раздел «Вёрстка» Приложения 1 ТЗ (T-28).

Семь проверок класса `file`: всё, что нужно для ответа, лежит в
`RenderedPresentation` и `TemplateSchema`. Модели не вызываются, результат
воспроизводим — это та часть аудита, которую можно предъявить как
доказуемо детерминированную.

**Три похожие проверки разведены по смыслу**, иначе один дефект порождал бы
три сообщения и пользователь чинил бы одно и то же трижды:

- `out_of_bounds` — про рамку элемента и край холста;
- `text_overflow` — про текст и его собственную рамку;
- `text_clipped` — про текст, который край холста действительно срезал:
  в рамку он помещался, в её видимую часть — уже нет.

**Вместимость считается метрикой вёрстки**, а не своей: `text_fits` из
`dpd.layout.overflow`. Своя метрика разошлась бы с той, по которой вёрстка
принимала решения, и аудит объявлял бы дефектом её норму.

**Правила берутся из шаблона, а не из головы.** Направляющие восстанавливаются
кластеризацией краёв слотов макета, поля — квантилью отступов слотов по всему
шаблону. Собственных значений «правильного» отступа у нас нет и быть не
может: шаблон чужой, и мерить его нашей линейкой означало бы выдавать свой
вкус за его правила.

**Шаблон не обвиняется в собственных решениях.** Содержимое, лежащее ровно
там, куда его поместил шаблон, нарушением полей не считается: в двух
калибровочных шаблонах из трёх есть слоты вплотную к краю холста, и находка
о них была бы претензией к автору шаблона, а не к нашей вёрстке.
"""

from __future__ import annotations

from itertools import combinations
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE

from dpd.audit.fixer import elements_of, fixes, slide_of, slot_bounds
from dpd.audit.package import read_package
from dpd.audit.registry import CheckSpec, check, param
from dpd.layout.overflow import text_fits
from dpd.models import (
    Bounds,
    Finding,
    RenderedElement,
    RenderedPresentation,
    TemplateSchema,
)

DEFAULT_SIZE_PT = 14.0
"""Кегль, если прогон его не несёт: то же значение, что принимает вёрстка."""

OUT_OF_BOUNDS = CheckSpec(
    id="layout.out_of_bounds",
    category="layout",
    check_class="file",
    severity="critical",
    fixability="mechanical",
    sublayer="4b",
    title="Элемент вышел за границы слайда",
)

OVERLAP = CheckSpec(
    id="layout.overlap",
    category="layout",
    check_class="file",
    severity="critical",
    fixability="lossy",
    sublayer="4b",
    title="Два блока наложились друг на друга",
)

TEXT_OVERFLOW = CheckSpec(
    id="layout.text_overflow",
    category="layout",
    check_class="file",
    severity="critical",
    fixability="lossy",
    sublayer="4b",
    title="Текст не поместился в свою рамку",
)

TEXT_CLIPPED = CheckSpec(
    id="layout.text_clipped",
    category="layout",
    check_class="file",
    severity="critical",
    fixability="mechanical",
    sublayer="4b",
    title="Текст обрезан краем слайда",
)

GUIDE_MISALIGN = CheckSpec(
    id="layout.guide_misalign",
    category="layout",
    check_class="file",
    severity="advice",
    fixability="mechanical",
    sublayer="4b",
    title="Блоки не выровнены по направляющим макета",
    plain="Блоки стоят неровно",
)

MARGIN_VIOLATION = CheckSpec(
    id="layout.margin_violation",
    category="layout",
    check_class="file",
    severity="warning",
    fixability="mechanical",
    sublayer="4b",
    title="Контент заходит в поля у краёв",
)

IMAGE_DISTORTED = CheckSpec(
    id="layout.image_distorted",
    category="layout",
    check_class="file",
    severity="warning",
    fixability="mechanical",
    sublayer="export",
    title="Картинка растянута, пропорции нарушены",
)


@check(OUT_OF_BOUNDS)
def check_out_of_bounds(deck: RenderedPresentation, tolerance: float | None = None) -> list[Finding]:
    """Найти элементы, вышедшие за край холста.

    Допуск нужен потому, что координаты — доли, пересчитанные из EMU: точное
    равенство единице здесь редкость, а находка о выходе на тысячную долю
    была бы шумом.
    """
    tolerance = _threshold(tolerance, OUT_OF_BOUNDS.id, "tolerance")

    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        for element in slide.elements:
            right = element.bounds.x + element.bounds.w - 1
            bottom = element.bounds.y + element.bounds.h - 1
            if max(right, bottom) <= tolerance:
                continue
            findings.append(
                OUT_OF_BOUNDS.finding(
                    f"Слайд {number}: блок «{element.slot_id}» выходит за холст "
                    f"{_direction(right, bottom)} и виден не целиком.",
                    slide_number=number,
                    slot_id=element.slot_id,
                    evidence={"right": round(right, 4), "bottom": round(bottom, 4)},
                )
            )
    return findings


@check(OVERLAP)
def check_overlap(deck: RenderedPresentation, area_tolerance: float | None = None) -> list[Finding]:
    """Найти пары блоков, наложившихся друг на друга.

    Допуск выражен долей площади меньшего блока, а не абсолютной величиной:
    касание соседних слотов после пересчёта долей даёт пересечение в сотые
    доли процента, и абсолютный порог пришлось бы подбирать под холст.
    """
    area_tolerance = _threshold(area_tolerance, OVERLAP.id, "area_tolerance")

    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        for first, second in combinations(slide.elements, 2):
            overlap = _intersection(first.bounds, second.bounds)
            smaller = min(_area(first.bounds), _area(second.bounds))
            if smaller <= 0 or overlap <= area_tolerance * smaller:
                continue
            findings.append(
                OVERLAP.finding(
                    f"Слайд {number}: блоки «{first.slot_id}» и «{second.slot_id}» "
                    f"наложились на {overlap / smaller:.0%} площади меньшего из них.",
                    slide_number=number,
                    slot_id=first.slot_id,
                    evidence={"share": round(overlap / smaller, 4), "with": second.slot_id},
                )
            )
    return findings


@check(TEXT_OVERFLOW)
def check_text_overflow(deck: RenderedPresentation) -> list[Finding]:
    """Найти текст, не поместившийся в свою рамку.

    Порога у проверки нет: вместимость считается метрикой гарнитуры и кегля,
    и настраивать здесь нечего. Вёрстка компенсирует переполнение кеглем, а
    когда шкала исчерпана — сокращает текст. **Сокращённый текст в рамку
    помещается, но не поместился**, и проверка называет его (T-60): иначе
    честная метрика превратила бы видимый на рендере хвост в тихую потерю
    содержания. Сработать без сокращения проверка должна там, где
    компенсации не справились.
    """
    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        cut = {item.slot_id: item for item in slide.applied_compensations if item.kind == "truncate"}
        for element in _text_elements(slide.elements):
            paragraphs, size = _paragraphs(element)
            fitted = text_fits(paragraphs, element.bounds, deck.canvas, size, element.frame)
            if fitted and element.slot_id in cut:
                lost = cut[element.slot_id]
                findings.append(
                    TEXT_OVERFLOW.finding(
                        f"Слайд {number}: текст блока «{element.slot_id}» не поместился и при "
                        f"наименьшем кегле шаблона {size:g} pt и сокращён "
                        f"с {lost.from_value:g} знаков до {lost.to_value:g}.",
                        slide_number=number,
                        slot_id=element.slot_id,
                        evidence={"sizePt": size, "truncated": True,
                                  "chars": lost.from_value, "kept": lost.to_value},
                    )
                )
                continue
            if fitted:
                continue
            findings.append(
                TEXT_OVERFLOW.finding(
                    f"Слайд {number}: текст блока «{element.slot_id}» не помещается "
                    f"в свою рамку при кегле {size:g} pt.",
                    slide_number=number,
                    slot_id=element.slot_id,
                    evidence={"sizePt": size, "chars": sum(len(item) for item in paragraphs)},
                )
            )
    return findings


@check(TEXT_CLIPPED)
def check_text_clipped(deck: RenderedPresentation, tolerance: float | None = None) -> list[Finding]:
    """Найти текст, срезанный краем холста.

    Срабатывает только тогда, когда текст помещался в свою рамку, а в её
    видимую часть — уже нет: именно это и означает «обрезан краем слайда».
    Текст, не влезающий и в полную рамку, — находка `layout.text_overflow`,
    и повторять её здесь значило бы заставить чинить один дефект дважды.
    """
    tolerance = _threshold(tolerance, TEXT_CLIPPED.id, "tolerance")

    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        for element in _text_elements(slide.elements):
            visible = _visible(element.bounds)
            if _area(element.bounds) - _area(visible) <= tolerance:
                continue

            paragraphs, size = _paragraphs(element)
            if not text_fits(paragraphs, element.bounds, deck.canvas, size, element.frame):
                continue
            if text_fits(paragraphs, visible, deck.canvas, size, element.frame):
                continue

            findings.append(
                TEXT_CLIPPED.finding(
                    f"Слайд {number}: край холста срезал часть текста блока "
                    f"«{element.slot_id}» — в видимую часть рамки он не помещается.",
                    slide_number=number,
                    slot_id=element.slot_id,
                    evidence={
                        "visibleShare": round(_area(visible) / _area(element.bounds), 4),
                        "sizePt": size,
                    },
                )
            )
    return findings


@check(GUIDE_MISALIGN)
def check_guide_misalign(
    deck: RenderedPresentation,
    template: TemplateSchema,
    tolerance: float | None = None,
    min_support: int | None = None,
    confidence: float | None = None,
) -> list[Finding]:
    """Найти блоки, не вставшие на направляющие макета.

    Направляющих в файле, как правило, нет: `ppt/viewProps.xml` есть лишь в
    одном калибровочном шаблоне из трёх, и там две из четырёх направляющих
    просто отмечают центр слайда. Поэтому они восстанавливаются по краям
    слотов макета: край, повторённый несколькими слотами, и есть линия, по
    которой макет выровнен.

    **Край собственного слота считается выровненным независимо от поддержки.**
    Иначе проверка обвиняет шаблон: на живых макетах у слота сплошь и рядом
    уникальный край, и блок, стоящий ровно там, куда его поместил автор
    шаблона, объявлялся бы невыровненным. Проверяется то, что сделали мы, —
    сдвиг блока с отведённого ему места мимо линий макета.

    Отсюда пониженная критичность и уверенность: правило выведено нами, а не
    прочитано у автора шаблона, и обвинять в его нарушении нельзя.
    """
    tolerance = _threshold(tolerance, GUIDE_MISALIGN.id, "tolerance")
    min_support = int(_threshold(min_support, GUIDE_MISALIGN.id, "min_support"))
    confidence = _threshold(confidence, GUIDE_MISALIGN.id, "confidence")

    layouts = {layout.id: layout for layout in template.layouts}

    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        layout = layouts.get(slide.layout_id)
        if layout is None or not layout.slots:
            continue

        slots = {slot.id: slot for slot in layout.slots}
        vertical = _guides([edge for s in layout.slots for edge in (s.bounds.x, s.bounds.x + s.bounds.w)], tolerance, min_support)
        horizontal = _guides([edge for s in layout.slots for edge in (s.bounds.y, s.bounds.y + s.bounds.h)], tolerance, min_support)

        for element in slide.elements:
            own = slots.get(element.slot_id)
            own_vertical = [own.bounds.x, own.bounds.x + own.bounds.w] if own else []
            own_horizontal = [own.bounds.y, own.bounds.y + own.bounds.h] if own else []

            off = []
            if vertical and not _on_guide(element.bounds.x, vertical + own_vertical, tolerance):
                off.append("вертикальной")
            if horizontal and not _on_guide(element.bounds.y, horizontal + own_horizontal, tolerance):
                off.append("горизонтальной")
            if not off:
                continue
            findings.append(
                GUIDE_MISALIGN.finding(
                    f"Слайд {number}: блок «{element.slot_id}» не выровнен по "
                    f"{' и '.join(off)} направляющей макета.",
                    slide_number=number,
                    slot_id=element.slot_id,
                    confidence=confidence,
                    evidence={
                        "x": round(element.bounds.x, 4),
                        "y": round(element.bounds.y, 4),
                        "verticalGuides": [round(value, 4) for value in vertical],
                        "horizontalGuides": [round(value, 4) for value in horizontal],
                    },
                )
            )
    return findings


@check(MARGIN_VIOLATION)
def check_margin_violation(
    deck: RenderedPresentation,
    template: TemplateSchema,
    quantile: float | None = None,
    confidence: float | None = None,
) -> list[Finding]:
    """Найти содержимое, зашедшее в поля у краёв холста.

    Поля выводятся из самого шаблона: `designTokens.spacing` в схеме нет, а
    придуманное значение мерило бы чужой шаблон нашей линейкой. Берётся
    квантиль отступов всех слотов — отступ, который шаблон выдерживает почти
    везде; минимум не годится, потому что одного слота вплотную к краю
    достаточно, чтобы правило исчезло.

    Находка выдаётся, только если элемент ближе к краю, чем **и** поле
    шаблона, **и** отведённое ему место. Содержимое, лежащее ровно в своём
    слоте, нарушением не считается — иначе мы предъявляли бы пользователю
    решения автора шаблона.
    """
    quantile = _threshold(quantile, MARGIN_VIOLATION.id, "quantile")
    confidence = _threshold(confidence, MARGIN_VIOLATION.id, "confidence")

    margins = _margins(template, quantile)
    if not any(margins.values()):
        return []

    slots = {(layout.id, s.id): s for layout in template.layouts for s in layout.slots}

    findings: list[Finding] = []
    for number, slide in enumerate(deck.slides, start=1):
        for element in slide.elements:
            own = slots.get((slide.layout_id, element.slot_id))
            gaps = _gaps(element.bounds)
            allowed = _gaps(own.bounds) if own is not None else None

            broken = {
                side: round(value, 4)
                for side, value in gaps.items()
                if value < margins[side] and (allowed is None or value < allowed[side])
            }
            if not broken:
                continue
            findings.append(
                MARGIN_VIOLATION.finding(
                    f"Слайд {number}: блок «{element.slot_id}» заходит в поля шаблона "
                    f"({', '.join(sorted(broken))}) ближе, чем отведённое ему место.",
                    slide_number=number,
                    slot_id=element.slot_id,
                    confidence=confidence,
                    evidence={"actual": broken, "expected": {k: round(v, 4) for k, v in margins.items()}},
                )
            )
    return findings


@check(IMAGE_DISTORTED)
def check_image_distorted(pptx_path: str | Path, aspect_tolerance: float | None = None) -> list[Finding]:
    """Найти картинки, поставленные не в своих пропорциях.

    Проверка работает по экспортированному файлу, а не по `RenderedPresentation`:
    изображений в контракте вёрстки нет вовсе — генерация изображений исключена
    решением D7, а картинки шаблона живут на макетах. Класс при этом остаётся
    `file`: соотношение сторон читается из разметки, а не из пикселей.

    Кадрирование учитывается: обрезанная картинка имеет другое исходное
    соотношение сторон, и без поправки на него каждая кадрированная картинка
    шаблона выглядела бы растянутой.
    """
    aspect_tolerance = _threshold(aspect_tolerance, IMAGE_DISTORTED.id, "aspect_tolerance")

    presentation, _ = read_package(pptx_path)
    if presentation is None:
        # Файл не открывается — об этом сообщает `integrity.file_broken`.
        return []

    findings: list[Finding] = []
    for number, slide in enumerate(presentation.slides, start=1):
        for shape in _pictures(slide.shapes):
            deviation = _aspect_deviation(shape)
            if deviation is None or deviation <= aspect_tolerance:
                continue
            findings.append(
                IMAGE_DISTORTED.finding(
                    f"Слайд {number}: картинка «{shape.name}» растянута — "
                    f"пропорции нарушены на {deviation:.0%}.",
                    slide_number=number,
                    evidence={"deviation": round(deviation, 4), "tolerance": aspect_tolerance},
                )
            )
    return findings


# --- Общее для проверок ---------------------------------------------------


def _threshold(value: float | None, check_id: str, name: str) -> float:
    """Порог из прогона, а при прямом вызове — из конфигурации."""
    return float(value) if value is not None else float(param(check_id, name))


def _direction(right: float, bottom: float) -> str:
    sides = []
    if right > 0:
        sides.append(f"вправо на {right:.1%}")
    if bottom > 0:
        sides.append(f"вниз на {bottom:.1%}")
    return " и ".join(sides) if sides else "незначительно"


def _area(bounds: Bounds) -> float:
    return bounds.w * bounds.h


def _intersection(first: Bounds, second: Bounds) -> float:
    width = min(first.x + first.w, second.x + second.w) - max(first.x, second.x)
    height = min(first.y + first.h, second.y + second.h) - max(first.y, second.y)
    return max(width, 0.0) * max(height, 0.0)


def _visible(bounds: Bounds) -> Bounds:
    """Часть рамки, оставшаяся внутри холста."""
    x, y = min(bounds.x, 1.0), min(bounds.y, 1.0)
    return Bounds(x=x, y=y, w=max(min(bounds.x + bounds.w, 1.0) - x, 0.0), h=max(min(bounds.y + bounds.h, 1.0) - y, 0.0))


def _text_elements(elements: list[RenderedElement]) -> list[RenderedElement]:
    return [element for element in elements if element.kind == "text" and element.runs]


def _paragraphs(element: RenderedElement) -> tuple[list[str], float]:
    """Абзацы элемента и кегль, которым они набраны.

    Берётся наибольший кегль прогонов: он определяет, поместится ли текст.
    """
    paragraphs = [run.text for run in element.runs]
    sizes = [run.size_pt for run in element.runs if run.size_pt]
    return paragraphs, max(sizes) if sizes else DEFAULT_SIZE_PT


def _guides(edges: list[float], tolerance: float, min_support: int) -> list[float]:
    """Восстановить направляющие: край, повторённый несколькими слотами."""
    guides: list[float] = []
    for edge in sorted(edges):
        cluster = [value for value in edges if abs(value - edge) <= tolerance]
        if len(cluster) < min_support:
            continue
        line = sum(cluster) / len(cluster)
        if not any(abs(line - existing) <= tolerance for existing in guides):
            guides.append(line)
    return guides


def _on_guide(value: float, guides: list[float], tolerance: float) -> bool:
    return any(abs(value - guide) <= tolerance for guide in guides)


def _gaps(bounds: Bounds) -> dict[str, float]:
    """Расстояния от рамки до каждого края холста."""
    return {
        "слева": bounds.x,
        "сверху": bounds.y,
        "справа": 1 - (bounds.x + bounds.w),
        "снизу": 1 - (bounds.y + bounds.h),
    }


def _margins(template: TemplateSchema, quantile: float) -> dict[str, float]:
    """Поля шаблона: отступ, который он выдерживает почти у всех слотов."""
    sides: dict[str, list[float]] = {"слева": [], "сверху": [], "справа": [], "снизу": []}
    for layout in template.layouts:
        for slot in layout.slots:
            for side, value in _gaps(slot.bounds).items():
                sides[side].append(max(value, 0.0))
    return {side: _quantile(values, quantile) for side, values in sides.items()}


def _quantile(values: list[float], share: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    return ordered[min(int(share * len(ordered)), len(ordered) - 1)]


def _pictures(shapes):
    """Картинки слайда, включая лежащие внутри групп."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pictures(shape.shapes)
        elif shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
            yield shape


def _aspect_deviation(shape) -> float | None:
    """На сколько соотношение сторон картинки отличается от исходного."""
    if not shape.width or not shape.height:
        return None
    try:
        native_width, native_height = shape.image.size
    except (AttributeError, ValueError):
        return None

    kept_width = native_width * (1 - shape.crop_left - shape.crop_right)
    kept_height = native_height * (1 - shape.crop_top - shape.crop_bottom)
    if kept_width <= 0 or kept_height <= 0:
        return None

    return abs((shape.width / shape.height) / (kept_width / kept_height) - 1)


# --- Исправления (T-33) ----------------------------------------------------
#
# Место, куда вернуть съехавший блок, известно из шаблона: это его слот. Если
# слот не найден или сам выходит за холст, блок прижимается к краю — хуже, чем
# по макету, но лучше, чем за кадром.


@fixes(OUT_OF_BOUNDS, touches="geometry")
def fix_out_of_bounds(deck, finding, template) -> bool:
    """Вернуть блок в холст: сначала на место по макету, иначе прижать к краю."""
    return _return_inside(deck, finding, template)


@fixes(TEXT_CLIPPED, touches="geometry")
def fix_text_clipped(deck, finding, template) -> bool:
    """Обрезка краем лечится тем же: блок должен оказаться внутри холста."""
    return _return_inside(deck, finding, template)


@fixes(MARGIN_VIOLATION, touches="geometry")
def fix_margin_violation(deck, finding, template) -> bool:
    """Отодвинуть блок из полей шаблона.

    Целевое значение берётся из доказательства находки: там записаны поля,
    которые шаблон выдерживает. Свои числа здесь брать неоткуда.
    """
    expected = finding.evidence.get("expected") or {}
    changed = False
    for element in elements_of(deck, finding):
        bounds = element.bounds
        x, y, w, h = bounds.x, bounds.y, bounds.w, bounds.h
        x = max(x, float(expected.get("слева", x)))
        y = max(y, float(expected.get("сверху", y)))
        right_gap = float(expected.get("справа", 1 - (x + w)))
        bottom_gap = float(expected.get("снизу", 1 - (y + h)))
        w = min(w, max(1 - right_gap - x, 0.0))
        h = min(h, max(1 - bottom_gap - y, 0.0))
        changed |= _place(element, x, y, w, h)
    return changed


@fixes(GUIDE_MISALIGN, touches="geometry")
def fix_guide_misalign(deck, finding, template) -> bool:
    """Притянуть блок к ближайшей направляющей макета.

    Направляющие взяты из доказательства: их восстановила сама проверка, и
    пересчитывать их здесь заново значило бы завести второй источник истины.
    """
    vertical = [float(value) for value in finding.evidence.get("verticalGuides", [])]
    horizontal = [float(value) for value in finding.evidence.get("horizontalGuides", [])]

    changed = False
    for element in elements_of(deck, finding):
        bounds = element.bounds
        x = min(vertical, key=lambda value: abs(value - bounds.x)) if vertical else bounds.x
        y = min(horizontal, key=lambda value: abs(value - bounds.y)) if horizontal else bounds.y
        changed |= _place(element, x, y, bounds.w, bounds.h)
    return changed


def _return_inside(deck, finding, template) -> bool:
    """Вернуть элементы находки внутрь холста."""
    slide = slide_of(deck, finding)
    if slide is None:
        return False

    changed = False
    for element in elements_of(deck, finding):
        place = slot_bounds(template, slide.layout_id, element.slot_id)
        if place is not None and place.x + place.w <= 1 and place.y + place.h <= 1:
            changed |= _place(element, place.x, place.y, place.w, place.h)
            continue

        bounds = element.bounds
        w, h = min(bounds.w, 1.0), min(bounds.h, 1.0)
        changed |= _place(element, min(bounds.x, 1 - w), min(bounds.y, 1 - h), w, h)
    return changed


def _place(element, x: float, y: float, w: float, h: float) -> bool:
    """Переставить элемент, сообщив, изменилось ли что-нибудь."""
    updated = Bounds(x=max(x, 0.0), y=max(y, 0.0), w=max(w, 0.0), h=max(h, 0.0))
    if updated == element.bounds:
        return False
    element.bounds = updated
    return True
